"""Make the reference templates offered beside each BoS document upload.

    python tools/make_bos_templates.py

Writes Word files into app/static/templates/, on the university letterhead
(app/doc_templates/jain_base.docx, through docgen._letterhead). They are
references for the department to fill, sign and upload; the Composition of
BoS Members follows the university's PAC form, row for row, so the portal's
row-by-row check reads it the same way.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from app.docgen import _borders, _letterhead  # noqa: E402

OUT = ROOT / "app" / "static" / "templates"
HEADER = "Board of Studies — Reference Template"


def _doc(title, sub=None):
    doc = _letterhead(HEADER, portrait=True)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    if sub:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.add_run(sub).italic = True
    return doc


def _details(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    _borders(t)
    for k in rows:
        c = t.add_row().cells
        c[0].text = k
        c[0].paragraphs[0].runs[0].bold = True
        c[1].text = ""
    _widths(t, [6.5, 10.5])
    doc.add_paragraph()
    return t


def _widths(t, cms):
    t.autofit = False
    for col, w in zip(t.columns, cms):
        col.width = Cm(w)
    for row in t.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)


def _table(doc, heads, n, first=None, widths=None):
    t = doc.add_table(rows=1, cols=len(heads))
    _borders(t)
    for i, h in enumerate(heads):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for k in range(n):
        cells = t.add_row().cells
        if first:
            cells[0].text = first(k)
    if widths is None and first is not None:
        # a narrow number column; the rest share the page
        rest = (17 - 1.4) / max(1, len(heads) - 1)
        widths = [1.4] + [rest] * (len(heads) - 1)
    if widths:
        _widths(t, widths)
    doc.add_paragraph()
    return t


def _signatures(doc, who):
    doc.add_paragraph()
    t = doc.add_table(rows=2, cols=len(who))
    for i, w in enumerate(who):
        t.rows[0].cells[i].text = "\n\n______________________"
        t.rows[1].cells[i].text = w
        t.rows[1].cells[i].paragraphs[0].runs[0].bold = True


def _retitle(p):
    """Re-word one paragraph, whose text Word may have split across runs."""
    full = p.text
    new = (full.replace("Program Assessment Committee (PAC)", "Board of Studies (BoS)")
               .replace("Program Assessment Committee", "Board of Studies")
               .replace("Role in PAC", "Role in BoS").replace("PAC", "BoS"))
    if new != full and p.runs:
        texts = [r for r in p.runs if r.text]
        texts[0].text = new
        for r in texts[1:]:
            r.text = ""


def composition():
    """The university's PAC form, re-titled for the Board of Studies."""
    doc = Document(str(OUT / "Composition_of_PAC.docx"))
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _retitle(p)
    for p in doc.paragraphs:
        _retitle(p)
    for sec in doc.sections:
        for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
            _retitle(p)
    doc.save(OUT / "Composition_of_BoS_Members.docx")


def external_profile():
    doc = _doc("Profile of External Expert", "One profile per external expert on the Board of Studies")
    _details(doc, ["Name", "Designation", "Organisation / Institution", "Highest qualification",
                   "Area of expertise", "Experience (years) — industry / academic",
                   "Role on the BoS (Industry Expert / Academician / Alumni)",
                   "E-mail", "Phone", "Address for correspondence"])
    doc.add_paragraph().add_run("Brief profile (achievements, publications, positions held)").bold = True
    _table(doc, ["Profile"], 1)
    _signatures(doc, ["Signature of the Expert", "Head of Department"])
    doc.save(OUT / "Profile_of_External_Expert.docx")


def minutes(title="Minutes of the Board of Studies Meeting", out="Minutes_of_BoS_Meeting.docx"):
    doc = _doc(title)
    _details(doc, ["Department", "School", "Academic year", "Date of the meeting", "Time",
                   "Venue / mode (in person, online)", "Chairperson"])
    doc.add_paragraph().add_run("Members present").bold = True
    _table(doc, ["S. No.", "Name", "Designation", "Category / Role in BoS"], 8, first=lambda k: str(k + 1))
    doc.add_paragraph().add_run("Agenda, discussion and resolutions").bold = True
    _table(doc, ["Item", "Agenda", "Discussion", "Resolution"], 6, first=lambda k: str(k + 1))
    doc.add_paragraph().add_run("Action taken / to be taken").bold = True
    _table(doc, ["S. No.", "Action", "Responsibility", "By when"], 4, first=lambda k: str(k + 1))
    _signatures(doc, ["Member Secretary", "Chairperson, BoS"])
    doc.save(OUT / out)


def attendance():
    doc = _doc("Attendance Sheet — Board of Studies Meeting")
    _details(doc, ["Department", "Date of the meeting", "Venue"])
    _table(doc, ["S. No.", "Name", "Designation & Organisation", "Category / Role in BoS", "Signature"],
           15, first=lambda k: str(k + 1))
    _signatures(doc, ["Member Secretary", "Chairperson, BoS"])
    doc.save(OUT / "Attendance_Sheet_BoS.docx")


def vision_mission():
    doc = _doc("Department Vision, Mission, PEOs and POs")
    _details(doc, ["Department", "School", "Programme(s)"])
    for head, rows, label in (("Vision", 1, None), ("Mission", 4, "M"),
                              ("Programme Educational Objectives (PEOs)", 4, "PEO"),
                              ("Programme Outcomes (POs)", 12, "PO"),
                              ("Programme Specific Outcomes (PSOs)", 3, "PSO")):
        doc.add_paragraph().add_run(head).bold = True
        if label:
            _table(doc, ["No.", "Statement"], rows, first=lambda k, lb=label: f"{lb}{k + 1}")
        else:
            _table(doc, ["Statement"], 1)
    doc.add_paragraph().add_run("Mapping of PEOs with the Mission (3 high · 2 medium · 1 low)").bold = True
    _table(doc, ["PEO / Mission", "M1", "M2", "M3", "M4"], 4, first=lambda k: f"PEO{k + 1}")
    _signatures(doc, ["Head of Department", "Dean / Director"])
    doc.save(OUT / "Vision_Mission_PEOs_POs.docx")


CRITERIA = [
    "The curriculum is relevant to current industry and societal needs",
    "Course objectives and outcomes are clearly stated",
    "The courses are sequenced appropriately across semesters",
    "Adequate weight is given to practical / hands-on learning",
    "Electives offer sufficient breadth and choice",
    "Emerging technologies and trends are covered",
    "The curriculum develops employability and soft skills",
    "Textbooks and references are current and adequate",
    "Assessment methods match the course outcomes",
    "Overall, the curriculum prepares students well",
]


def feedback(title, out, extra=None):
    doc = _doc(title, "Rate each statement: 5 Strongly agree · 4 Agree · 3 Neutral · 2 Disagree · 1 Strongly disagree")
    _details(doc, ["Name", "Category (Industry Expert / Academician / Alumni / Employer / Student / Parent)",
                   "Organisation", "Designation", "E-mail", "Programme reviewed", "Date"] + (extra or []))
    t = _table(doc, ["S. No.", "Statement", "5", "4", "3", "2", "1"], len(CRITERIA),
               first=lambda k: str(k + 1), widths=[1.2, 10.8, 1, 1, 1, 1, 1])
    for k, c in enumerate(CRITERIA):
        t.rows[k + 1].cells[1].text = c
    doc.add_paragraph().add_run("Suggestions for improvement").bold = True
    _table(doc, ["Suggestions"], 1)
    _signatures(doc, ["Signature of the Stakeholder", "Head of Department"])
    doc.save(OUT / out)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    composition()
    external_profile()
    minutes()
    minutes("Minutes of the Pre-Board of Studies Meeting", "Minutes_of_Pre-BoS_Meeting.docx")
    attendance()
    vision_mission()
    feedback("Stakeholder Feedback — Curriculum Design and Development",
             "Stakeholder_Feedback_Curriculum.docx")
    feedback("Stakeholder Feedback — Proposed New Programme", "Stakeholder_Feedback_New_Programme.docx",
             extra=["Proposed programme", "Is there demand for this programme? (Yes / No, with reasons)"])
    for p in sorted(OUT.glob("*.docx")):
        print(p.relative_to(ROOT))


if __name__ == "__main__":
    main()
