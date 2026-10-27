"""A syllabus sheet follows the University's sample: Course Information taken
from the Curriculum table (never typed again), Course Objectives, the Course
Outcomes each with its Bloom's Taxonomy level, the CO-PO-PSO Mapping with its
Articulation row — each course with its own PO and PSO columns, from PO1 and
PSO1 — the Course Content, and the Text Books, Book References and Web
References. Course Objectives start from the programme's Objective."""

import io

from test_workflow import login


def test_course_information_comes_from_the_curriculum_line():
    from app.schema import course_info
    r = {"semester": 1, "l": 2, "t": 0, "p": 2, "e": 3, "credits": 4, "cia": 70, "ese": 30}
    assert course_info(r) == {"semester": 1, "ltpe": "2-0-2-3", "hours_per_week": 7, "credits": 4,
                              "teaching_hours": 60, "ca_ese": "70:30"}
    # what the line does not have yet is left out, not guessed
    assert course_info({"credits": "3"}) == {"credits": 3, "teaching_hours": 45}
    assert course_info({}) == {}


def test_the_programmes_objective_is_read_line_by_line():
    from app.schema import objective_lines
    assert objective_lines("1. To build skills\n2) To prepare students\n\n- To develop ethics\nii) To lead") == \
        ["To build skills", "To prepare students", "To develop ethics", "To lead"]
    assert objective_lines("") == [] and objective_lines(None) == []


def test_course_objectives_follow_the_programmes_until_the_course_writes_its_own():
    from app.workflow import complete_syllabus, course_objectives
    c = complete_syllabus({"course_title": "Accounting", "outcomes": "Apply"}, "To build skills")
    assert c["objectives"] == "To build skills"
    assert course_objectives(c, "To build skills\nTo serve") == "To build skills\nTo serve"   # it follows
    assert course_objectives({"objectives": ""}, "To serve") == "To serve"
    c["objectives"] = "To keep books of account"                                        # its own
    assert course_objectives(c, "To serve") == "To keep books of account"
    assert complete_syllabus(dict(c), "To serve")["objectives"] == "To keep books of account"
    # a programme with no Objective yet: the course's own, or a start to edit
    assert complete_syllabus({"course_title": "Accounting"})["objectives"].startswith("To explain")


def test_an_outcomes_opening_verb_suggests_its_level():
    from app.schema import bloom_level
    assert bloom_level("Apply fundamental accounting concepts") == 3
    assert bloom_level("Evaluate the choices.") == 5
    assert bloom_level("Design a ledger") == 6
    assert bloom_level("Be aware of things") is None
    assert bloom_level("") is None


def test_each_course_has_its_own_po_and_pso_columns():
    from app.schema import course_columns
    assert course_columns({}) == ["PO1", "PSO1"]                       # a new course starts with one of each
    assert course_columns({"po_count": 3, "pso_count": 2}) == ["PO1", "PO2", "PO3", "PSO1", "PSO2"]
    # a sheet that never said: as many as its mapping holds
    cols = course_columns({"co_po": [{"PO1": "3", "PO9": "2", "PSO4": "1"}]})
    assert len(cols) == 13 and cols[8] == "PO9" and cols[-1] == "PSO4"
    assert course_columns({"po_count": "99", "pso_count": "0"}) == [f"PO{n}" for n in range(1, 16)] + ["PSO1"]
    # the rule reads the course's own columns only
    from app.validation import r_co_po_mapping_valid as rule
    assert rule(_sheet(po_count=1, pso_count=1, co_po=[{"PO1": "3", "PO5": "9"}, {"PSO1": "2"}]), {}, "courses") == []


def _sheet(**kw):
    c = {"course_code": "26BCM1C01", "course_title": "Financial Accounting",
         "co_table": [{"text": "Apply accounting concepts", "bloom": "3"},
                      {"text": "Prepare final accounts", "bloom": "3"}],
         "co_po": [{"PO1": "3", "PSO1": "2"}, {"PO2": "2"}]}
    c.update(kw)
    return {"courses": [c]}


def test_every_course_outcome_needs_its_level():
    from app.validation import r_outcomes_have_bloom_levels as rule
    assert rule(_sheet(), {}, "courses") == []
    issues = rule(_sheet(co_table=[{"text": "Apply concepts", "bloom": ""}]), {}, "courses")
    assert [i["level"] for i in issues] == ["error"] and "CO 1" in issues[0]["message"]
    # an older sheet, with its outcomes only as lines, is asked for the levels
    old = _sheet(outcomes="Apply\nPrepare")
    del old["courses"][0]["co_table"]
    assert rule(old, {}, "courses")[0]["level"] == "error"


def test_the_mapping_holds_a_number_in_each_cell():
    from app.validation import r_co_po_mapping_valid as rule
    assert rule(_sheet(), {}, "courses") == []
    # any number is taken, as typed
    assert rule(_sheet(co_po=[{"PO1": "5"}, {"PO2": "2.5"}]), {}, "courses") == []
    issues = rule(_sheet(co_po=[{"PO1": "high"}, {"PO2": "2"}]), {}, "courses")
    assert any(i["level"] == "error" and "must be a number" in i["message"] for i in issues)
    # a CO left unmapped is a warning; a course with no mapping at all is an error
    issues = rule(_sheet(co_po=[{"PO1": "3"}, {"PO1": "-"}]), {}, "courses")
    assert [i["level"] for i in issues] == ["warning"] and "CO2" in issues[0]["message"]
    issues = rule(_sheet(co_po=[]), {}, "courses")
    assert any(i["level"] == "error" and "fill in its CO-PO-PSO Mapping" in i["message"] for i in issues)


def test_the_page_and_the_word_syllabus_follow_the_sample(app, client):
    from docx import Document

    from app import demo_dept
    from app.db import get_db
    with app.app_context():
        user, pw = demo_dept.create(app.config["ACADEMIC_YEAR"], app.config["UPLOAD_ROOT"])
    login(client, user, pw)
    client.post("/department/demo/fill-all")
    items = client.get("/department/api/record?programme=DEMOBBA").get_json()["items"]

    # the Curriculum no longer asks how many POs and PSOs: the mapping has the sample's
    cur = items[0]["data"]
    assert "Number of POs" not in client.get("/department/stage/prog_curriculum/DEMOBBA").get_data(as_text=True)
    page = client.get("/department/stage/prog_syllabus/DEMOBBA").get_data(as_text=True)
    assert 'outcome_max: {"po": 15, "pso": 8}' in page
    # the programme's Objective is what each course's Course Objectives start from
    cur["profile"]["objective"] = "1. To build sound knowledge\n2. To prepare students for careers"
    assert client.post("/department/api/prog_curriculum/DEMOBBA/save", json=cur).get_json()["ok"]
    page = client.get("/department/stage/prog_syllabus/DEMOBBA").get_data(as_text=True)
    assert 'programme_objectives: "To build sound knowledge\\nTo prepare students for careers"' in page
    assert all(c["objectives"].startswith("To give students of the") for c in items[1]["data"]["courses"])
    # each course's information is there already, from its line of the table
    first = next(r for r in cur["semester_structure"] if r.get("course_code") and "/" not in r["course_code"])
    assert f'"ltpe": "{first["l"]}-{first["t"]}-{first["p"]}-{first["e"]}"' in page.replace(".0", "")

    # the sample sheets pass the rules
    data = items[1]["data"]
    j = client.post("/department/api/prog_syllabus/DEMOBBA/validate", json=data).get_json()
    assert not [i for i in j["issues"] if i["level"] == "error"], j["issues"]
    # every course takes PO1-PO3 and PSO1-PSO2: the Word syllabus maps to those
    for c in data["courses"]:
        c.update(po_count=3, pso_count=2)
    assert client.post("/department/api/prog_syllabus/DEMOBBA/save", json=data).get_json()["ok"]

    from app.docgen import syllabus_docx
    with app.app_context():
        doc = Document(io.BytesIO(syllabus_docx(get_db(), "DEMO", app.config["ACADEMIC_YEAR"], "DEMOBBA").getvalue()))
    text = "\n".join(p.text for p in doc.paragraphs)
    # the sample courses' objectives were never edited: they follow the programme's
    assert "1.  To build sound knowledge\n2.  To prepare students for careers\nCourse Outcomes (COs):" in text
    for heading in ("Course Information:", "Course Objectives:", "Course Outcomes (COs):",
                    "CO-PO-PSO Mapping:", "Course Content:", "Text Books:", "Book References:",
                    "3-High, 2-Medium, 1-Low"):
        assert heading in text, heading
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    info = next(t for t in tables if t[0][0] == "Course Title")
    assert [row[0] for row in info] == ["Course Title", "Course Code", "Semester", "Credits"]
    assert [row[2] for row in info[1:]] == ["L-T-P-E", "Contact Hours", "CA: ESE"]
    outcomes = next(t for t in tables if t[0][0] == "Course Outcome")
    assert outcomes[0] == ["Course Outcome", "Description", "Bloom’s Taxonomy Level"]
    assert outcomes[1][0] == "CO 1" and outcomes[1][2] in {"1", "2", "3", "4", "5", "6"}
    mapping = next(t for t in tables if t[0][0] == "CO")
    assert mapping[0] == ["CO", "PO1", "PO2", "PO3", "PSO1", "PSO2"]
    assert mapping[-1][0] == "Articulation" and len(mapping) == len(outcomes) + 1
