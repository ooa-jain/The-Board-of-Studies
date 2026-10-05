"""The Office's analysis sheets, comments asking for a change, and the
curriculum's link from each course to its syllabus."""

from test_workflow import login


def _demo_filled_and_submitted(app, client):
    from app import demo_dept
    with app.app_context():
        user, pw = demo_dept.create("2027-28", app.config["UPLOAD_ROOT"])
    login(client, user, pw)
    client.post("/department/demo/fill-all")
    assert client.post("/department/api/submit-all").get_json()["ok"]
    client.get("/logout")
    return user, pw


def test_sheets_show_every_department_and_download(app, client):
    _demo_filled_and_submitted(app, client)
    login(client, app.config["ADMIN_USERNAME"], app.config["ADMIN_PASSWORD"])
    for url in ("/admin/submissions", "/admin/analysis", "/admin/sheets?tab=programmes",
                "/admin/sheets?tab=documents", "/admin/sheets?tab=comments"):
        r = client.get(url)
        assert r.status_code == 200, url
        assert 'class="ov-tabs"' in r.get_data(as_text=True), "one Overview, one row of tabs"
    assert "Department of Demonstration Studies" in client.get("/admin/submissions").get_data(as_text=True)
    body = client.get("/admin/sheets?tab=programmes").get_data(as_text=True)
    assert "BBA in Demonstration Management" in body and ">38<" in body
    r = client.get("/admin/sheets.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Stages", "Programmes", "Documents", "Comments"]


def test_a_comment_reopens_a_stage_and_the_department_answers_it(app, client):
    from app.db import get_db
    user, pw = _demo_filled_and_submitted(app, client)
    login(client, app.config["ADMIN_USERNAME"], app.config["ADMIN_PASSWORD"])
    client.post("/admin/comments", data={"dept": "DEMO", "stage": "bos_documents",
                                         "text": "Upload the signed minutes.", "reopen": "on"})
    with app.app_context():
        c = get_db().comments.find_one({"dept_code": "DEMO"})
        assert c["reopened"]
        assert get_db().submissions.find_one({"dept_code": "DEMO"})["stages"]["bos_documents"]["status"] == "returned"
    review = client.get("/admin/submissions/DEMO/review/bos_documents").get_json()
    assert review["comments"][0]["text"] == "Upload the signed minutes."

    client.get("/logout")
    login(client, user, pw)
    assert "Upload the signed minutes." in client.get("/department/").get_data(as_text=True)
    page = client.get("/department/stage/bos_documents").get_data(as_text=True)
    assert "Comment from the Office of Academics" in page
    client.post(f"/department/comments/{c['_id']}/done", data={"reply": "Uploaded the signed copy."})
    with app.app_context():
        c = get_db().comments.find_one({"_id": c["_id"]})
    assert c["status"] == "done" and c["reply"] == "Uploaded the signed copy."


def test_curriculum_rows_link_to_their_syllabus(app, client):
    from app import demo_dept
    with app.app_context():
        user, pw = demo_dept.create("2027-28", app.config["UPLOAD_ROOT"])
    login(client, user, pw)
    client.post("/department/demo/fill-all")
    for key in ("dept_info", "pre_bos", "bos_documents"):
        from app.db import get_db
        with app.app_context():
            data = get_db().submissions.find_one({"dept_code": "DEMO"})["stages"][key]["data"]
        assert client.post(f"/department/api/{key}/submit", json=data).get_json()["ok"]
    page = client.get("/department/stage/prog_curriculum/DEMOBBA").get_data(as_text=True)
    syl = page.split("syllabi: ")[1].split(",\n")[0]
    import json
    syllabi = json.loads(syl)
    current = next(s for s in syllabi if s["current"])
    assert current["url"].endswith("/department/stage/prog_syllabus/DEMOBBA")
    assert {"code": "27BBA1C01", "title": "Principles of Management"} in current["courses"]
    assert any(not s["current"] for s in syllabi), "the earlier batches are there too"
