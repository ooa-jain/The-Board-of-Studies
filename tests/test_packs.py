"""The CS & IT Drive pack: the data in it, loading it, what admin and department see."""

import re

from test_workflow import login, make_department

MCA = ("MCAREG", "MCACYS", "MCASCT", "MCAISM", "MCAIML")


def _admin(app, client):
    login(client, app.config["ADMIN_USERNAME"], app.config["ADMIN_PASSWORD"])


def _load(app):
    from app.packs import load_pack
    with app.app_context():
        return load_pack("csit", app.config["ACADEMIC_YEAR"], "test")


def _sub(app, code):
    from app.db import get_db
    with app.app_context():
        return get_db().submissions.find_one({"dept_code": code})


# ---------------------------------------------------------------- the pack

def _structure(code):
    from app.packs import get_pack
    return get_pack("csit")["programmes"][code]["prog_curriculum"]["semester_structure"]


def test_pack_matrices_add_up_to_the_printed_semester_totals():
    for code in MCA:
        by_sem = {}
        for r in _structure(code):
            by_sem[r["semester"]] = by_sem.get(r["semester"], 0) + r["credits"]
        assert by_sem == {1: 23, 2: 23, 3: 23, 4: 21}, code


def test_every_pack_course_reconciles_l_t_p_e_with_credits():
    from app.validation import RULES
    for code in MCA:
        data = {"semester_structure": _structure(code)}
        issues = RULES["ltpe_credit_arithmetic"](data, {}, "semester_structure")
        assert not [i for i in issues if "works out to" in i["message"]], code


def test_pack_course_codes_are_unique_and_well_formed():
    from app.schema import STAGE_BY_KEY
    from app.validation import RULES
    sec = next(s for s in STAGE_BY_KEY["prog_curriculum"]["sections"]
               if s["key"] == "semester_structure")
    pattern = next(c["pattern"] for c in sec["columns"] if c["name"] == "course_code")
    for code in MCA:
        data = {"semester_structure": _structure(code)}
        assert not RULES["unique_course_codes"](data, {}, "semester_structure"), code
        for r in data["semester_structure"]:
            assert re.match(pattern, r["course_code"]), r["course_code"]


def test_pack_programmes_are_the_departments_catalogue_programmes():
    from app import catalogue
    from app.packs import get_pack
    pack = get_pack("csit")
    official = {r["programme_code"] for r in catalogue.load()
                if catalogue._name(r["department"]) == catalogue._name(pack["department"]["dept_name"])}
    assert set(pack["programmes"]) <= official


# ---------------------------------------------------------------- loading

def test_loading_fills_the_existing_department_rather_than_adding_one(app):
    make_department(app, "CSI-JYN", "Department of Computer Science and IT",
                    "School of Computer Science & Information Technology")
    report = _load(app)
    assert not report["created"] and report["dept"]["dept_code"] == "CSI-JYN"
    assert report["password"] is None                 # it already had a login
    sub = _sub(app, "CSI-JYN")
    assert set(sub["programmes"]) == set(MCA) | {"BCAGAI"}


def test_loading_creates_the_department_and_a_login_when_missing(app):
    report = _load(app)
    assert report["created"] and report["password"]
    sub = _sub(app, report["dept"]["dept_code"])
    reg = sub["programmes"]["MCAREG"]
    data = reg["prog_curriculum"]["data"]
    assert len(data["semester_structure"]) == 36
    assert data["details"]["degree_level"] == "PG - 2 Year"
    dsa = next(c for c in reg["prog_syllabus"]["data"]["courses"] if c["course_code"] == "25MCAC101")
    assert dsa["modules"][0] == {"title": "List ADT", "hours": 10, "revised": dsa["modules"][0]["revised"]}
    assert len(dsa["outcomes"].splitlines()) == 5


def test_loading_validates_like_a_real_submit(app):
    report = _load(app)
    curr = _sub(app, report["dept"]["dept_code"])["programmes"]["MCAREG"]["prog_curriculum"]
    assert curr["status"] == "draft"
    missing = {i["field"] for i in curr["issues"] if i["level"] == "error"}
    # not in the Drive folder; everything else comes from the matrix or the standard wording
    assert {"intake", "eligibility"} <= missing


def test_loading_again_keeps_what_is_already_there(app):
    from app.db import get_db
    report = _load(app)
    code = report["dept"]["dept_code"]
    with app.app_context():
        get_db().submissions.update_one(
            {"dept_code": code},
            {"$set": {"programmes.MCAREG.prog_curriculum.data.profile.intake": 60}})
    again = _load(app)
    assert again["password"] is None and not again["written"]
    data = _sub(app, code)["programmes"]["MCAREG"]["prog_curriculum"]["data"]
    assert data["profile"]["intake"] == 60


def test_the_department_can_open_its_filled_curriculum_and_syllabus(app, client):
    report = _load(app)
    login(client, report["dept"]["username"], report["password"])
    for part in ("prog_curriculum", "prog_syllabus"):
        r = client.get(f"/department/stage/{part}/MCAREG", follow_redirects=True)
        assert r.status_code == 200
        assert "25MCAC101" in r.get_data(as_text=True)


# ---------------------------------------------------------------- the rules it needed

def test_an_elective_pair_is_two_courses_for_the_syllabus(app):
    from app.workflow import course_fill_source
    report = _load(app)
    courses = course_fill_source(_sub(app, report["dept"]["dept_code"]), "MCAREG")
    by_code = {c["course_code"]: c["course_title"] for c in courses}
    assert by_code["25MCAGE2041"] == "Essentials of Cloud Computing"
    assert by_code["25MCAGE2042"] == "Essentials of Cyber Security"


def test_a_code_reused_inside_an_elective_pair_is_caught():
    from app.validation import RULES
    data = {"s": [{"course_code": "25MCA101"}, {"course_code": "25MCA102 / 25MCA101"}]}
    assert any("25MCA101" in i["message"] for i in RULES["unique_course_codes"](data, {}, "s"))


def test_experiential_hours_count_one_credit_per_three():
    from app.validation import RULES
    project = {"s": [{"course_code": "P1", "l": 2, "t": 0, "p": 4, "e": 15, "credits": 9}]}
    assert not [i for i in RULES["ltpe_credit_arithmetic"](project, {}, "s")
                if "works out to" in i["message"]]


# ---------------------------------------------------------------- admin

def test_admin_loads_the_pack_from_the_departments_page(app, client):
    make_department(app, "CSI-JYN", "Department of Computer Science and IT",
                    "School of Computer Science & Information Technology")
    _admin(app, client)
    assert "Load from Drive" in client.get("/admin/departments").get_data(as_text=True)
    r = client.post("/admin/packs/csit/load")
    assert r.status_code == 302 and r.headers["Location"].endswith("/admin/submissions/CSI-JYN")
    assert "Load again" in client.get("/admin/departments").get_data(as_text=True)
    page = client.get("/admin/submissions/CSI-JYN").get_data(as_text=True)
    assert "Filled from the department" in page and "TBA-" in page


# ---------------------------------------------------------------- add a programme from Curriculum

def _signed_in_department(app, client):
    report = _load(app)
    login(client, report["dept"]["username"], report["password"])
    client.post("/department/who", data={"name": "Test Faculty", "email": "t@example.edu",
                                         "next": "/department/"})
    return report["dept"]["dept_code"]


def test_curriculum_offers_to_add_a_ug_or_pg_programme(app, client):
    _signed_in_department(app, client)
    page = client.get("/department/stage/curriculum").get_data(as_text=True)
    assert "Add a UG programme" in page and "Add a PG programme" in page
    assert "/department/programmes/add?level=PG" in page


def test_adding_a_programme_reopens_department_information_and_comes_back(app, client):
    from app.db import get_db
    code = _signed_in_department(app, client)
    with app.app_context():                      # as if Department Information were submitted
        get_db().submissions.update_one({"dept_code": code},
                                        {"$set": {"stages.dept_info.status": "submitted"}})
    r = client.get("/department/programmes/add?level=UG")
    assert r.status_code == 302
    assert r.headers["Location"].endswith("/department/stage/dept_info?add=UG&back=curriculum")
    assert _sub(app, code)["stages"]["dept_info"]["status"] == "draft"

    page = client.get(r.headers["Location"]).get_data(as_text=True)
    assert "Back to Curriculum" in page and "Adding a new" in page
    assert "back=curriculum" in page             # submitting returns to Curriculum


def test_a_sealed_record_cannot_be_reopened_from_curriculum(app, client):
    from app.db import get_db
    code = _signed_in_department(app, client)
    with app.app_context():
        get_db().submissions.update_one({"dept_code": code}, {"$set": {"status": "sealed"}})
    r = client.get("/department/programmes/add?level=PG")
    assert r.headers["Location"].endswith("/department/stage/curriculum")


# ---------------------------------------------------------------- the Drive files

def test_loading_attaches_the_drive_files_where_the_department_would_upload_them(app):
    from pathlib import Path
    report = _load(app)
    code = report["dept"]["dept_code"]
    sub = _sub(app, code)
    assert {"MCAREG": "2026-2028_MCA GEN_Course MAtrix.xlsx"}["MCAREG"] == \
        sub["programmes"]["MCAREG"]["prog_curriculum"]["data"]["curriculum_file"]["document"]["name"]
    syl = sub["programmes"]["MCASCT"]["prog_syllabus"]["data"]["syllabus_file"]["document"]
    assert syl["name"] == "MCA_SCT_SYLLABUS_2026_28.docx"
    stored = Path(app.config["UPLOAD_ROOT"]) / app.config["ACADEMIC_YEAR"] / code / \
        "prog_syllabus" / syl["stored"]
    assert stored.stat().st_size == syl["size"] == 619194
    # kept with the record: the checklist and the articulation sheets
    names = [f["name"] for f in sub["imported_from"]["files"]]
    assert any(n.startswith("Checklist for Document Submission") for n in names)
    assert "2026-2028_MCA-ISMS-Articulation.xlsx" in names


def test_loading_again_does_not_attach_the_files_twice(app):
    from app.db import get_db
    report = _load(app)
    code = report["dept"]["dept_code"]
    with app.app_context():
        before = get_db().files.count_documents({"dept_code": code})
    again = _load(app)
    assert not again["attached"]
    with app.app_context():
        assert get_db().files.count_documents({"dept_code": code}) == before


def test_attached_files_show_in_the_admins_uploaded_documents(app, client):
    report = _load(app)
    _admin(app, client)
    page = client.get(f"/admin/submissions/{report['dept']['dept_code']}").get_data(as_text=True)
    assert "MCA_ISMS_SYLLABUS_2026_28.docx" in page
    from app.db import get_db
    with app.app_context():
        rec = get_db().files.find_one({"original_name": "MCA_ISMS_SYLLABUS_2026_28.docx"})
    r = client.get(f"/admin/documents/{rec['dept_code']}/{rec['stored_name']}")
    assert r.status_code == 200 and len(r.data) == 3582763


def test_the_drive_strip_shows_the_departments_login(app, client):
    report = _load(app)
    _admin(app, client)
    page = client.get("/admin/departments").get_data(as_text=True)
    assert report["dept"]["username"] in page
    assert report["password"] in page and "Load again" in page


# ---------------------------------------------------------------- the single-sheet report

def test_department_excel_is_one_sheet_with_ticks_and_links(app, client):
    import io

    from openpyxl import load_workbook
    report = _load(app)
    code = report["dept"]["dept_code"]
    _admin(app, client)
    r = client.get(f"/admin/export/{code}.xlsx")
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Report"]
    ws = wb["Report"]
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    text = " ".join(str(c.value) for c in cells)
    assert "Department information" in text and "Pre-BoS" in text
    assert "PG programmes" in text and "✗" in text
    links = [c.hyperlink.target for c in cells if c.hyperlink]
    assert any("/share/r/" in u for u in links)                  # programme reports
    assert any("/share/d/" in u for u in links)                  # the attached Drive files


def test_department_word_is_one_report_with_links(app, client):
    import io

    from docx import Document
    report = _load(app)
    _admin(app, client)
    r = client.get(f"/admin/export/{report['dept']['dept_code']}.docx")
    doc = Document(io.BytesIO(r.data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "UG programmes" in text and "PG programmes" in text
    rels = [rel.target_ref for rel in doc.part.rels.values() if "hyperlink" in rel.reltype]
    assert any("/share/r/" in u for u in rels)


def test_the_programme_report_opens_in_the_browser(app, client):
    report = _load(app)
    _admin(app, client)
    page = client.get(f"/admin/report/{report['dept']['dept_code']}/MCAREG").get_data(as_text=True)
    assert "Programme structure" in page and "25MCAC101" in page
    assert "Data Structures" in page and "Total: 90 credits" in page


def test_a_departments_own_download_links_to_its_own_pages(app, client):
    import io

    from openpyxl import load_workbook
    from app.db import get_db
    code = _signed_in_department(app, client)
    with app.app_context():          # the download is the finished record
        get_db().submissions.update_one({"dept_code": code}, {"$set": {"status": "sealed"}})
    r = client.get("/department/export.xlsx")
    ws = load_workbook(io.BytesIO(r.data))["Report"]
    links = [c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink]
    assert links and all("/share/" in u for u in links)



# ---------------------------------------------------------------- share links (no sign-in)

def _download_links(app, client, code):
    import io

    from openpyxl import load_workbook
    _admin(app, client)
    ws = load_workbook(io.BytesIO(client.get(f"/admin/export/{code}.xlsx").data))["Report"]
    return {c.value: c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink}


def test_a_document_link_opens_without_signing_in(app, client):
    from urllib.parse import urlparse
    code = _load(app)["dept"]["dept_code"]
    links = _download_links(app, client, code)
    stranger = app.test_client()                         # not signed in
    docx = urlparse(links["MCA_ISMS_SYLLABUS_2026_28.docx"]).path
    page = stranger.get(docx)
    assert page.status_code == 200
    body = page.get_data(as_text=True)
    assert "MCA_ISMS_SYLLABUS_2026_28.docx" in body and "DATA STRUCTURES" in body.upper()
    raw = stranger.get(docx + "/file?download=1")
    assert raw.status_code == 200 and len(raw.data) == 3582763
    sheet = stranger.get(urlparse(links["2026-2028_MCA GEN_Course MAtrix.xlsx"]).path)
    assert "25MCAC101" in sheet.get_data(as_text=True)


def test_a_programme_report_link_opens_without_signing_in(app, client):
    from urllib.parse import urlparse
    code = _load(app)["dept"]["dept_code"]
    links = _download_links(app, client, code)
    report = next(u for name, u in links.items() if name == "Master of Computer Applications")
    page = app.test_client().get(urlparse(report).path).get_data(as_text=True)
    assert "Programme structure" in page and "25MCAC101" in page
    assert "/admin/" not in page                         # its document links are share links too


def test_a_changed_share_link_does_not_open(app):
    c = app.test_client()
    assert c.get("/share/d/not-a-real-code").status_code == 404
    assert c.get("/share/r/WyJyIiwiWCJd.tampered").status_code == 404
