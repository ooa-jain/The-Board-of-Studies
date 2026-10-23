"""An earlier batch is three pages, as the current batch is — its Curriculum,
Syllabus and Course Revision — in only the semesters Settings asks it for;
and a programme's whole record, every batch, downloads as one ZIP."""

import io
import zipfile

from test_workflow import login


def _semesters(app, by_year):
    from app.db import get_db
    with app.app_context():
        get_db().settings.update_one({"_id": "app"}, {"$set": {
            "current_batch": "2026-2027", "existing_batches": ["2024-2025", "2025-2026"],
            "batch_semesters": by_year}}, upsert=True)


def _demo(app, client):
    from app import demo_dept
    with app.app_context():
        user, pw = demo_dept.create(app.config["ACADEMIC_YEAR"], app.config["UPLOAD_ROOT"])
    login(client, user, pw)
    client.post("/department/demo/fill-all")
    return "DEMOBBA"


def test_an_earlier_batch_revision_starts_from_its_own_syllabus():
    from app.workflow import batch_revision_rows
    data = {"courses": [
        {"course_code": "24BBA5C01", "course_title": "Strategy", "semester": 5,
         "modules": [{"title": "Basics", "hours": 9, "revised": "What a strategy is"}]},
        {"course_code": "24BBA3C01", "course_title": "Costing", "semester": 3, "modules": []},
    ]}
    rows = batch_revision_rows(data, "2024–2025", "2026", [5, 6])
    assert [r["course_code"] for r in rows] == ["24BBA5C01"]           # only the batch's semesters
    r = rows[0]
    assert r["year_previous"] == "2024" and r["year_latest"] == "2026" and r["prev_code"] == "24BBA5C01"
    assert r["modules"] == [{"previous": "Basics (9 Hrs) — What a strategy is",
                             "revised": "Basics (9 Hrs) — What a strategy is"}]


def test_an_earlier_batch_offers_only_its_semesters(app, client):
    _semesters(app, {"2024": [5, 6], "2025": [3, 4]})
    code = _demo(app, client)
    page = client.get(f"/department/stage/prog_syllabus_b2024/{code}").get_data(as_text=True)
    assert "revision_courses" in page and "Semesters 5 and 6" in page       # its three pages, its semesters
    assert '"batch_semesters": [5, 6]' in page
    assert "Sem 5–6" in page and "Sem 3–4" in page                       # the side menu and the batch switch
    # the side menu heads the current batch's steps with its year, as it does the earlier batches
    assert "Current batch · 2026–2027" in page and "Earlier batches" in page
    # its sample data sits in those semesters, and has its Course Revision
    from app.db import get_db
    with app.app_context():
        data = get_db().submissions.find_one({"dept_code": "DEMO"})["programmes"][code]["prog_syllabus_b2024"]["data"]
    assert data["batch_courses"] and {int(r["semester"]) for r in data["batch_courses"]} <= {5, 6}
    assert data["revision_courses"] and {int(r["semester"]) for r in data["revision_courses"]} <= {5, 6}
    # a course outside them is refused
    bad = dict(data, batch_courses=[dict(data["batch_courses"][0], semester=1)])
    j = client.post(f"/department/api/prog_syllabus_b2024/{code}/validate", json=bad).get_json()
    assert any("semester 1 is not one of this batch's" in i["message"] for i in j["issues"])
    j = client.post(f"/department/api/prog_syllabus_b2024/{code}/validate", json=data).get_json()
    assert not [i for i in j["issues"] if i["level"] == "error"], j["issues"]


def test_a_programme_downloads_every_batch_in_one_zip(app, client):
    _semesters(app, {"2024": [5, 6], "2025": [3, 4]})
    code = _demo(app, client)
    j = client.get(f"/department/api/record?programme={code}").get_json()
    assert j["zip"] == f"/department/programme/{code}/report.zip"
    r = client.get(j["zip"])
    assert r.status_code == 200 and r.mimetype == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    folders = sorted({n.split("/")[1] for n in names})
    assert folders == ["1 Current batch 2026-2027", "2 Earlier batch 2024-2025 (Sem 5-6)",
                       "3 Earlier batch 2025-2026 (Sem 3-4)"]
    for f in folders:
        for doc in ("Curriculum", "Syllabus", "Course Revision"):
            assert any(n.endswith(f"/{f}/{doc}.docx") for n in names), (f, doc)
    # an earlier batch's documents say which batch and semesters, and hold its courses
    from docx import Document
    z = zipfile.ZipFile(io.BytesIO(r.data))
    text = lambda n: "\n".join(p.text for p in Document(io.BytesIO(z.read(n))).paragraphs) + " ".join(
        c.text for t in Document(io.BytesIO(z.read(n))).tables for row in t.rows for c in row.cells)
    cur = next(n for n in names if "2 Earlier batch" in n and n.endswith("Curriculum.docx"))
    assert "Curriculum — Batch of 2024–2025 · Semesters 5 and 6" in text(cur) and "24" in text(cur)
    rev = next(n for n in names if "2 Earlier batch" in n and n.endswith("Course Revision.docx"))
    assert "Batch of 2024–2025" in text(rev)
    # a submitted part offers it beside Review
    client.post(f"/department/api/prog_syllabus_b2024/{code}/submit",
                json=client.get(f"/department/api/record?programme={code}").get_json()["items"][3]["data"])
    page = client.get(f"/department/stage/prog_syllabus_b2024/{code}").get_data(as_text=True)
    assert f'href="/department/programme/{code}/report.zip" download' in page
    # another programme's, or one that is not there, is not given
    assert client.get("/department/programme/NOPE/report.zip").status_code == 404

    # the Office gets the same, and the department's ZIP holds every programme's
    client.get("/logout")
    login(client, app.config["ADMIN_USERNAME"], app.config["ADMIN_PASSWORD"])
    page = client.get(f"/admin/report/DEMO/{code}").get_data(as_text=True)
    assert "Download all (ZIP)" in page and "Earlier batch 2024–2025 · Sem 5–6" in page
    r = client.get(f"/admin/report/DEMO/{code}/report.zip")
    assert r.status_code == 200 and len(zipfile.ZipFile(io.BytesIO(r.data)).namelist()) == 9
    r = client.get("/admin/export/DEMO.zip")
    names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
    assert any(n.startswith(f"Programmes/{code} - ") and n.endswith("Course Revision.docx") for n in names)
