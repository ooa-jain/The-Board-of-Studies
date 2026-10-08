"""
A programme's Curriculum and Syllabus as Word documents, laid out as the
Office of Academics' own templates are:

    Curriculum   the BCom (Corporate Finance) programme-structure document —
                 landscape, the JAIN logo and the programme's line in the
                 header, "Prepared and Approved by Office of Academics" in the
                 footer; the title block, the profile (items 1 to 12), the
                 Classification of Credits, the Programme Structure semester
                 by semester, the Summary and Annexure I (minors)
    Syllabus     Template_Syllabus — one bordered sheet per course: the
                 programme, code and name, credits and hours, pedagogy,
                 course outcomes, modules with their hours, skill development
                 activities and books for reference

Both are built from what the department entered, so they are always current.
The letterhead (logo, header line, footer) comes from
app/doc_templates/jain_base.docx — the curriculum template with its body
removed.
"""

from __future__ import annotations

import copy
import io
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

from .schema import DEFAULT_REGULATIONS, STAGE_BY_KEY
from .workflow import programmes_of

BASE = Path(__file__).parent / "doc_templates" / "jain_base.docx"

# the template's own colours
GREEN = "C6E0B4"      # table titles and column heads
PALE = "D9EAD3"       # course-group bands
BLUE = "ADD8E6"       # Honours / Honours with Research bands
NAVY = "1F4E79"       # semester bands, headings

GROUP_HEADS = [
    ("Major (Core)", "MAJOR COURSES"),
    ("Discipline Specific Elective (DSE)", "DISCIPLINE SPECIFIC ELECTIVES"),
    ("Minor Stream", "MINOR COURSES"),
    ("Multidisciplinary", "MULTIDISCIPLINARY / OPEN ELECTIVE COURSES"),
    ("Ability Enhancement Courses (AEC)", "ABILITY ENHANCEMENT COURSES"),
    ("Skill Enhancement Courses (SEC)", "SKILL ENHANCEMENT COURSES"),
    ("Value Added Courses (VAC)", "VALUE ADDED COURSES"),
    ("Summer Internship", "SUMMER INTERNSHIP"),
    ("Research Project / Dissertation", "RESEARCH PROJECT / DISSERTATION"),
    ("Mandatory Non-Credit Course", "MANDATORY NON-CREDIT COURSES"),
    ("Mandatory Non-Credit Audit Course", "MANDATORY NON-CREDIT AUDIT COURSES"),
]

# Classification of Credits columns, as the template heads them
DIST_GROUPS = [
    ("Major (Core)", "Major"),
    ("Minor Stream", "Minor"),
    ("Multidisciplinary", "Multi-Disciplinary / OE"),
    ("Ability Enhancement Courses (AEC)", "Ability Enhancement / AEC"),
    ("Skill Enhancement Courses (SEC)", "Skill Enhancement / SEC"),
    ("Value Added Courses (VAC)", "Common Value Added / VAC"),
    ("Summer Internship", "Summer Internship / INTERNSHIP"),
    ("Research Project / Dissertation", "Research Project / Dissertation / PROJECT"),
]
NC_COURSE = "Mandatory Non-Credit Course"
NC_AUDIT = "Mandatory Non-Credit Audit Course"


# ---------------------------------------------------------------------------
# small Word helpers
# ---------------------------------------------------------------------------

def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else f


def _fmt(v):
    v = _num(v)
    return "" if v is None else str(v)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(table, size=6):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tbl_pr.append(borders)


def _widths(table, widths):
    """Fixed column widths (EMU), on the grid as well as each cell — Word
    reads the cells, LibreOffice the grid."""
    table.autofit = False
    widths = [Emu(int(w)) for w in widths]
    grid = table._tbl.tblGrid
    for i, col in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            col.set(qn("w:w"), str(int(widths[i].twips)))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for row in table.rows:
        seen = set()
        for i, cell in enumerate(row.cells):
            if id(cell._tc) in seen or i >= len(widths):
                continue
            seen.add(id(cell._tc))
            span = int(cell._tc.grid_span or 1)
            cell.width = Emu(sum(int(w) for w in widths[i:i + span]))


def _write(cell, text, bold=False, size=9, align=None, color=None, italic=False):
    """Replace a cell's text with one or more lines."""
    cell.text = ""
    lines = str(text if text is not None else "").split("\n")
    p = cell.paragraphs[0]
    for n, line in enumerate(lines):
        if n:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        if align is not None:
            p.alignment = align
        run = p.add_run(line)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return cell


def _labelled(cell, label, value, size=10, align=None):
    """ "Label: value" with the label in bold, as the syllabus sheet prints it."""
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(size)
    if value not in (None, ""):
        r2 = p.add_run(" " + str(value))
        r2.font.size = Pt(size)
    return p


def _band(table, text, fill, color="000000", bold=True, size=9):
    row = table.add_row()
    cell = row.cells[0].merge(row.cells[-1])
    _write(cell, text, bold=bold, size=size, align=WD_ALIGN_PARAGRAPH.CENTER, color=color)
    _shade(cell, fill)
    return row


def _para(doc, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=None, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _lines(text):
    """A textarea's entries, one per line, without any numbering typed in."""
    out = []
    for line in str(text or "").splitlines():
        line = re.sub(r"^\s*(?:[a-zA-Z]\)|\(?[0-9ivx]+[.)]|[-•*])\s+", "", line).strip()
        if line:
            out.append(line)
    return out


# ---------------------------------------------------------------------------
# the letterhead
# ---------------------------------------------------------------------------

def _letterhead(header_line, portrait=False):
    doc = Document(str(BASE))
    sec = doc.sections[0]
    if portrait:
        # the template is landscape; turn it and pull the logo and the red
        # rule in by the width the page loses
        w, h = sec.page_width, sec.page_height
        shrink = int(w) - int(h)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = h, w
        sec.left_margin = sec.right_margin = Emu(720000)
        for el in sec.header._element.iter():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "posOffset" and el.text and int(el.text) > 5_000_000:
                el.text = str(int(el.text) - shrink)
            elif tag in ("extent", "ext") and int(el.get("cx", 0)) > 5_000_000:
                el.set("cx", str(int(el.get("cx")) - shrink))
    # the programme's own line where the template names BCom
    for p in sec.header.paragraphs:
        if p.text.strip():
            # runs with no text hold the logo and the rule: leave them be
            texts = [r for r in p.runs if r.text]
            texts[0].text = header_line
            for r in texts[1:]:
                r.text = ""
            break
    if portrait:
        # the footer pads "Prepared and Approved…" out to the page number for
        # the landscape width; on a portrait page that pushes it to a new line
        for p in sec.footer.paragraphs:
            for r in p.runs:
                if r.text.strip().startswith("Prepared"):
                    r.text = r.text.rstrip() + " " * max(1, len(r.text) - len(r.text.rstrip()) - 115)
    for st in doc.styles:
        if st.name == "Normal":
            st.font.size = Pt(10)
    return doc


def _title(text):
    """Word keeps a document title to 255 characters."""
    text = str(text)
    return text if len(text) <= 255 else text[:254] + "…"


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# a programme's stored record
# ---------------------------------------------------------------------------

def programme_record(db, dept_code, year, programme_code):
    """(dept, programme, curriculum data, syllabus courses) or None."""
    dept = db.departments.find_one({"dept_code": dept_code})
    if not dept:
        return None
    sub = db.submissions.find_one({"dept_code": dept_code, "academic_year": year}) or {}
    prog = next((p for p in programmes_of(sub, dept)
                 if p["programme_code"].upper() == str(programme_code).upper()), None)
    if not prog:
        return None
    stored = (sub.get("programmes") or {}).get(prog["programme_code"]) or {}
    cur = (stored.get("prog_curriculum") or {}).get("data") or {}
    syl = ((stored.get("prog_syllabus") or {}).get("data") or {}).get("courses") or []
    return dept, prog, cur, [c for c in syl if isinstance(c, dict)]


def _title_line(prog, cur):
    details = cur.get("details") or {}
    name = details.get("programme_name") or prog.get("programme_name") or prog["programme_code"]
    return name, (details.get("specialisation") or prog.get("specialisation") or ""), \
        (details.get("batch") or prog.get("batch") or "")


def _version_line(name, batch):
    return f"{name} {batch}_v1_{datetime.now().strftime('%d/%m/%Y')}".replace("  ", " ")


# ---------------------------------------------------------------------------
# Curriculum
# ---------------------------------------------------------------------------

def _blocks(rows):
    """The template's three runs of semesters: the common ones, then the
    Honours and the Honours with Research semesters (7 and 8)."""
    track = lambda r: r.get("track") or "All semesters"    
    special = {_num(r.get("semester")) for r in rows if track(r) != "All semesters"}
    base = [r for r in rows if _num(r.get("semester")) not in special]
    blocks = [(None, base)]
    for t in ("Honours", "Honours with Research"):
        rs = [r for r in rows if _num(r.get("semester")) in special
              and track(r) in (t, "All semesters")]
        if any(track(r) == t for r in rs):
            blocks.append((t, rs))
    return base, blocks


def _sems(rows):
    return sorted({_num(r.get("semester")) for r in rows if _num(r.get("semester")) is not None})


def _in_group(r, cat):
    c = r.get("nep_category")
    return c == cat or (cat == "Major (Core)" and c == "Discipline Specific Elective (DSE)")


def _credits(rows):
    return sum(_num(r.get("credits")) or 0 for r in rows)


def _profile_table(doc, cur, width, degree_level=None):
    from .schema import PG_REGULATIONS, for_level, is_pg
    prof = cur.get("profile") or {}
    details = cur.get("details") or {}
    t = doc.add_table(rows=0, cols=3)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = next(s for s in for_level(STAGE_BY_KEY["prog_curriculum"], degree_level)["sections"]
                  if s["key"] == "profile")["fields"]
    defaults = {**DEFAULT_REGULATIONS, **(PG_REGULATIONS if is_pg(degree_level) else {})}
    for n, f in enumerate(fields, start=1):
        label = re.sub(r"^\d+\.\s*", "", f["label"])
        row = t.add_row().cells
        _write(row[0], str(n), align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
        _write(row[1], label, size=9.5)
        value = prof.get(f["name"])
        if f["name"] == "reservation_policy":
            cell = row[2]
            cell.text = ""
            for part in f.get("fixed_table") or []:
                _labelled(cell, part["head"] + ":", "  ".join(part["items"]), size=9)
                p = cell.add_paragraph(part.get("note") or "")
                p.runs[0].font.size = Pt(8.5) if p.runs else None
            continue
        if f["name"] == "duration_months":
            m = _num(value)
            value = f"{_fmt(m / 12)} years / {_fmt(m)} months" if m else ""
        if f["name"] == "course_specialisation" and not value:
            value = " — ".join(x for x in (details.get("programme_name"),
                                          details.get("specialisation")) if x)
        if value in (None, "") and f["name"] in defaults:
            value = defaults[f["name"]]
        if f["name"] == "objective":
            items = _lines(value)
            value = "\n".join(f"{i}. {x}" for i, x in enumerate(items, start=1))
        _write(row[2], _fmt(value) if isinstance(value, (int, float)) else (value or ""), size=9.5)
    _widths(t, [width * 0.06, width * 0.17, width * 0.77])


def _classification_table(doc, rows, width):
    base, blocks = _blocks(rows)
    heads = (["Semester"] + [g[1] for g in DIST_GROUPS] +
             ["Total Credits", "No. of Mandatory Non-Credit Course/s",
              "No. of Mandatory Non-Credit Audit Course/s"])
    t = doc.add_table(rows=0, cols=len(heads))
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _band(t, "Classification of Credits and Number of Non-Credit Courses", GREEN, size=10)
    hr = t.add_row().cells
    for i, h in enumerate(heads):
        _write(hr[i], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER,
               color="FF0000" if "Audit" in h else None)
        _shade(hr[i], GREEN)

    nc = lambda rs: len([r for r in rs if r.get("nep_category") == NC_COURSE or      
                         (_num(r.get("credits")) == 0 and r.get("nep_category") != NC_AUDIT)]) or ""
    audit = lambda rs: len([r for r in rs if r.get("nep_category") == NC_AUDIT]) or ""

    def line(label, rs, bold=False):
        cells = t.add_row().cells
        vals = ([label] + [_fmt(_credits([r for r in rs if _in_group(r, g[0])])) for g in DIST_GROUPS]
                + [_fmt(_credits(rs)), str(nc(rs)), str(audit(rs))])
        for i, v in enumerate(vals):
            _write(cells[i], v, bold=bold, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    for n, (title, rs) in enumerate(blocks):
        if title:
            _band(t, title, BLUE)
        for s in _sems(rs):
            line(str(s), [r for r in rs if _num(r.get("semester")) == s])
        line("TOTAL", rs if n == 0 else base + rs, bold=True)
    w = width / len(heads)
    _widths(t, [w] * len(heads))


def _structure_table(doc, rows, width):
    t = doc.add_table(rows=0, cols=7)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _band(t, "Programme Structure", GREEN)
    heads = ["Course Code", "Course Title", "Major / Minor", "Credits",
             "Continuous Assessment Marks", "Term End Examination Marks", "Total Marks"]
    hr = t.add_row().cells
    for i, h in enumerate(heads):
        _write(hr[i], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hr[i], GREEN)

    order = {c: n for n, (c, _) in enumerate(GROUP_HEADS)}
    heads_by = dict(GROUP_HEADS)
    track = lambda r: r.get("track") or "All semesters"    
    for s in _sems(rows):
        in_sem = [r for r in rows if _num(r.get("semester")) == s]
        tracks = ["All semesters"] + [x for x in ("Honours", "Honours with Research")
                                      if any(track(r) == x for r in in_sem)]
        for tr in tracks:
            rs = [r for r in in_sem if track(r) == tr]
            if not rs:
                continue
            title = f"SEMESTER {s}" + ("" if tr == "All semesters" else f" — {tr.upper()}")
            _band(t, title, NAVY, color="FFFFFF")
            groups = sorted({r.get("nep_category") or "" for r in rs}, key=lambda c: order.get(c, 99))
            for g in groups:
                _band(t, heads_by.get(g, (g or "OTHER COURSES").upper()), PALE)
                for r in [x for x in rs if (x.get("nep_category") or "") == g]:
                    mm = ("Major" if _in_group(r, "Major (Core)") else
                          "Minor" if g == "Minor Stream" else "")
                    cells = t.add_row().cells
                    vals = [r.get("course_code") or "", r.get("course_title") or "", mm,
                            _fmt(r.get("credits")), _fmt(r.get("cia")), _fmt(r.get("ese")),
                            _fmt(r.get("total_marks") or ((_num(r.get("cia")) or 0)
                                                          + (_num(r.get("ese")) or 0)) or "")]
                    for i, v in enumerate(vals):
                        _write(cells[i], v, size=9,
                               align=None if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER)
            tot = t.add_row().cells
            _write(tot[0], "Total", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            _write(tot[3], _fmt(_credits(rs)), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            _write(tot[6], _fmt(sum(_num(r.get("total_marks")) or 0 for r in rs)), bold=True,
                   size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _widths(t, [width * x for x in (0.12, 0.40, 0.08, 0.07, 0.11, 0.11, 0.11)])


def _summary_table(doc, rows, width, title="SUMMARY"):
    base, blocks = _blocks(rows)
    t = doc.add_table(rows=0, cols=5)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _band(t, title, GREEN, size=10)
    heads = ["Semester", "100% Continuous Assessment Credits",
             "Term End (University) Examination Credits", "Total Credits", "Total Marks"]
    hr = t.add_row().cells
    for i, h in enumerate(heads):
        _write(hr[i], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    ca_only = lambda r: _num(r.get("ese")) == 0 and _num(r.get("cia")) is not None

    def line(label, rs, bold=False):
        cells = t.add_row().cells
        ca = _credits([r for r in rs if ca_only(r)])
        vals = [label, _fmt(ca) if ca else "-", _fmt(_credits([r for r in rs if not ca_only(r)])),
                _fmt(_credits(rs)), _fmt(sum(_num(r.get("total_marks")) or 0 for r in rs))]
        for i, v in enumerate(vals):
            _write(cells[i], v, bold=bold, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    for n, (title, rs) in enumerate(blocks):
        if title:
            _band(t, title, "FFFFFF")
        for s in _sems(rs):
            line(str(s), [r for r in rs if _num(r.get("semester")) == s])
        line("Total Credits", rs if n == 0 else base + rs, bold=True)
    _widths(t, [width * 0.13, width * 0.17, width * 0.18, width * 0.11, width * 0.11])


def _minors_table(doc, minors, width):
    t = doc.add_table(rows=0, cols=5)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _band(t, "Annexure – I – Minor", GREEN, size=10)
    hr = t.add_row().cells
    for i, h in enumerate(("Minor", "Semester", "Course Code", "Course Title", "Credits")):
        _write(hr[i], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hr[i], GREEN)
    for stream in dict.fromkeys(m.get("minor_title") or "" for m in minors):
        rs = sorted([m for m in minors if (m.get("minor_title") or "") == stream],
                    key=lambda m: _num(m.get("semester")) or 0)
        _band(t, stream or "Minor", PALE)
        for m in rs:
            cells = t.add_row().cells
            for i, v in enumerate(("", _fmt(m.get("semester")), m.get("course_code") or "",
                                   m.get("course_title") or "", _fmt(m.get("credits")))):
                _write(cells[i], v, size=9,
                       align=None if i in (2, 3) else WD_ALIGN_PARAGRAPH.CENTER)
    _widths(t, [width * 0.18, width * 0.09, width * 0.14, width * 0.45, width * 0.09])


PG_SUMMARY = [
    ("Generic Core", "Generic Core"),
    ("Generic Elective", "Generic Elective"),
    ("Specialisation Core", "Specialisation Core"),
    ("Specialisation Elective", "Specialisation Elective"),
    ("Open Elective", "Open Elective"),
    ("Research / Thesis / Project / Patent", "RESEARCH / THESIS / PROJECT / PATENT"),
]


def _pg_summary_table(doc, rows, width):
    """The PG Course Matrix's SUMMARY: credits by classification, by semester."""
    heads = ["Semester"] + [h for _, h in PG_SUMMARY] + ["Mandatory Non-Credit", "TOTAL CREDITS"]
    t = doc.add_table(rows=0, cols=len(heads))
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _band(t, "SUMMARY", GREEN, size=10)
    hr = t.add_row().cells
    for i, h in enumerate(heads):
        _write(hr[i], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hr[i], GREEN)
    nc = lambda rs: len([r for r in rs if r.get("nep_category") == NC_COURSE]) or ""

    def line(label, rs, bold=False):
        cells = t.add_row().cells
        vals = ([label] + [_fmt(_credits([r for r in rs if r.get("nep_category") == g])) for g, _ in PG_SUMMARY]
                + [str(nc(rs)), _fmt(_credits(rs))])
        for i, v in enumerate(vals):
            _write(cells[i], v, bold=bold, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    for s in _sems(rows):
        line(str(s), [r for r in rows if _num(r.get("semester")) == s])
    line("TOTAL", rows, bold=True)
    _widths(t, [width / len(heads)] * len(heads))


def _pg_semester_tables(doc, rows, width):
    """One table a semester: Course Code, Course Title, Category, Credits and
    the marks — as the PG Course Matrix lays them out."""
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
    order = {c: n for n, c in enumerate([g for g, _ in PG_SUMMARY] + [NC_COURSE])}
    for s in _sems(rows):
        rs = sorted([r for r in rows if _num(r.get("semester")) == s],
                    key=lambda r: (order.get(r.get("nep_category"), 99), str(r.get("course_code") or "")))
        t = doc.add_table(rows=0, cols=7)
        _borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        head = t.add_row()
        # the head row repeats if the semester runs onto the next page
        head._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        hr = head.cells
        for i, h in enumerate(("Course Code", "Course Title", "Category", "Credits",
                               "Continuous Assessment Marks", "Term End Examination Marks", "Total Marks")):
            _write(hr[i], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade(hr[i], GREEN)
        _band(t, f"SEMESTER {roman[int(s) - 1] if 0 < int(s) <= 8 else s}", NAVY, color="FFFFFF")
        for r in rs:
            cells = t.add_row().cells
            vals = [r.get("course_code") or "", r.get("course_title") or "", r.get("nep_category") or "",
                    _fmt(r.get("credits")), _fmt(r.get("cia")), _fmt(r.get("ese")),
                    _fmt(r.get("total_marks") or ((_num(r.get("cia")) or 0) + (_num(r.get("ese")) or 0)) or "")]
            for i, v in enumerate(vals):
                _write(cells[i], v, size=9, align=None if i in (0, 1, 2) else WD_ALIGN_PARAGRAPH.CENTER)
        tot = t.add_row().cells
        _write(tot[0], "Total", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(tot[3], _fmt(_credits(rs)), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(tot[6], _fmt(sum(_num(r.get("total_marks")) or 0 for r in rs)), bold=True,
               size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _widths(t, [width * x for x in (0.12, 0.36, 0.16, 0.07, 0.10, 0.10, 0.09)])
        doc.add_paragraph()


def curriculum_docx(db, dept_code, year, programme_code):
    rec = programme_record(db, dept_code, year, programme_code)
    if not rec:
        return None
    _dept, prog, cur, _ = rec
    from .schema import is_pg, pg_groups
    deg = (cur.get("details") or {}).get("degree_level") or prog.get("degree_level")
    if is_pg(deg):
        return _pg_curriculum_docx(prog, pg_groups(cur, deg), deg)
    name, spec, batch = _title_line(prog, cur)
    doc = _letterhead(_version_line(name, batch))
    sec = doc.sections[0]
    width = int(sec.page_width - sec.left_margin - sec.right_margin)

    _para(doc, "JAIN (Deemed-to-be University), Bangalore", bold=True, size=15)
    _para(doc, name, bold=True, size=14)
    if spec:
        _para(doc, spec, size=14)
    _para(doc, f"Programme Structure {batch}".strip(), bold=True, size=14, after=6)
    _profile_table(doc, cur, width)

    rows = [r for r in (cur.get("semester_structure") or [])
            if isinstance(r, dict) and _num(r.get("semester")) is not None]
    _page_break(doc)
    if rows:
        _classification_table(doc, rows, width)
    else:
        _para(doc, "Classification of Credits — the programme structure is not entered yet.",
              size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

    _page_break(doc)
    _para(doc, "Programme Structure", bold=True, size=13, color=NAVY, after=6)
    if rows:
        _structure_table(doc, rows, width)
        doc.add_paragraph()
        _summary_table(doc, rows, width)
    else:
        _para(doc, "No courses entered yet.", size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

    minors = [m for m in (cur.get("minors") or []) if isinstance(m, dict)]
    if minors:
        _page_break(doc)
        _minors_table(doc, minors, width)
    doc.core_properties.title = _title(f"{name} — Curriculum {batch}".strip())
    return _save(doc)


def _pg_curriculum_docx(prog, cur, deg):
    """A PG programme's curriculum in the PG Course Matrix's layout: the
    profile (1–13), the SUMMARY by classification, a table a semester, and
    the credits by assessment."""
    name, spec, batch = _title_line(prog, cur)
    doc = _letterhead(_version_line(name, batch))
    sec = doc.sections[0]
    width = int(sec.page_width - sec.left_margin - sec.right_margin)
    _para(doc, "Jain (Deemed-to-be University), Bangalore", bold=True, size=15)
    _para(doc, name, bold=True, size=14)
    if spec:
        _para(doc, spec, size=14)
    _para(doc, f"Programme Structure {batch}".strip(), bold=True, size=14, after=6)
    _profile_table(doc, cur, width, deg)

    rows = [r for r in (cur.get("semester_structure") or [])
            if isinstance(r, dict) and _num(r.get("semester")) is not None]
    _page_break(doc)
    _para(doc, "Jain (Deemed-to-be University), Bangalore", bold=True, size=13)
    _para(doc, name, bold=True, size=12)
    _para(doc, f"Programme Structure {batch}".strip(), bold=True, size=12, after=6)
    if rows:
        _pg_summary_table(doc, rows, width)
        _para(doc, "The classifications defined for assigning credits: Generic Core, Generic Elective, "
                   "Specialisation Core, Specialisation Elective, Open Elective, and Research / Thesis / "
                   "Project / Patent. Electives vary by programme; each elective a student may choose is "
                   "listed in its semester.", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        _page_break(doc)
        _pg_semester_tables(doc, rows, width)
        _summary_table(doc, rows, width, title="CREDITS BY ASSESSMENT")
    else:
        _para(doc, "No courses entered yet.", size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.core_properties.title = _title(f"{name} — Programme Structure {batch}".strip())
    return _save(doc)


# ---------------------------------------------------------------------------
# Syllabus
# ---------------------------------------------------------------------------

def _course_sheet(doc, programme_name, c, width):
    t = doc.add_table(rows=0, cols=3)
    _borders(t, size=8)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    center = WD_ALIGN_PARAGRAPH.CENTER

    def full():
        row = t.add_row()
        cell = row.cells[0].merge(row.cells[2])
        cell.text = ""
        return cell

    head = full()
    _labelled(head, "Name of the Program:", programme_name, align=center)
    _labelled(head, "Course Code:", c.get("course_code") or "", align=center)
    _labelled(head, "Name of the Course:", "", align=center).add_run(
        " " + (c.get("course_title") or "")).bold = True

    hr = t.add_row().cells
    for i, h in enumerate(("Course Credits", "No. of Hours per Week", "Total No. of Teaching Hours")):
        _write(hr[i], h, bold=True, size=10, align=center)
    vr = t.add_row().cells
    for i, v in enumerate((f"{_fmt(c.get('credits'))} Credits" if _fmt(c.get("credits")) else "",
                           f"{_fmt(c.get('hours_per_week'))} Hrs" if _fmt(c.get("hours_per_week")) else "",
                           f"{_fmt(c.get('teaching_hours'))} Hrs" if _fmt(c.get("teaching_hours")) else "")):
        _write(vr[i], v, bold=True, size=10, align=center)

    ped = full()
    _labelled(ped, "Pedagogy:", c.get("pedagogy") or
              "Classrooms lecture, Case studies, Tutorial Classes, Group discussion, Seminar & field work etc.,")

    co = full()
    _labelled(co, "Course Outcomes: On successful completion of the course, the students' will be able to",
              "")
    for i, line in enumerate(_lines(c.get("outcomes"))):
        p = co.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_after = Pt(0)
        p.add_run(f"{chr(97 + i % 26)})  {line}").font.size = Pt(10)

    row = t.add_row()
    left = row.cells[0].merge(row.cells[1])
    _write(left, "Syllabus:", bold=True, size=10)
    _write(row.cells[2], "Hours", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for n, m in enumerate([m for m in (c.get("modules") or []) if isinstance(m, dict)], start=1):
        row = t.add_row()
        left = row.cells[0].merge(row.cells[1])
        _write(left, f"Module No. {n}: {m.get('title') or ''}".rstrip(": "), bold=True, size=10)
        _write(row.cells[2], _fmt(m.get("hours")), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
        body = str(m.get("revised") or m.get("content") or "").strip()
        if body:
            cell = full()
            _write(cell, body, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def numbered(label, text, note=None):
        cell = full()
        _labelled(cell, label, "")
        for i, line in enumerate(_lines(text), start=1):
            p = cell.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(0)
            p.add_run(f"{i}.  {line}").font.size = Pt(10)
        if note:
            p = cell.add_paragraph()
            r = p.add_run(note)
            r.bold = True
            r.font.size = Pt(10)

    numbered("Skill Development Activities:", c.get("skill_activities"))
    numbered("Books for reference:", c.get("books"), "Note: Latest edition of books may be used.")
    _widths(t, [width * 0.30, width * 0.36, width * 0.34])


def syllabus_docx(db, dept_code, year, programme_code):
    rec = programme_record(db, dept_code, year, programme_code)
    if not rec:
        return None
    _dept, prog, cur, courses = rec
    name, spec, batch = _title_line(prog, cur)
    doc = _letterhead(_version_line(name, batch), portrait=True)
    sec = doc.sections[0]
    width = int(sec.page_width - sec.left_margin - sec.right_margin)

    _para(doc, "JAIN (Deemed-to-be University), Bangalore", bold=True, size=14)
    _para(doc, name, bold=True, size=13)
    if spec:
        _para(doc, spec, size=12)
    _para(doc, f"Syllabus {batch}".strip(), bold=True, size=13, after=8)
    if not courses:
        _para(doc, "No syllabus entered yet.", size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    courses = sorted(courses, key=lambda c: (_num(c.get("semester")) or 99))
    sem = object()
    for n, c in enumerate(courses):
        if n:
            _page_break(doc)
        if c.get("semester") != sem and _num(c.get("semester")):
            sem = c.get("semester")
            _para(doc, f"SEMESTER {_fmt(sem)}", bold=True, size=11, color=NAVY, after=4)
        _course_sheet(doc, name, copy.deepcopy(c), width)
    doc.core_properties.title = _title(f"{name} — Syllabus {batch}".strip())
    return _save(doc)


# ---------------------------------------------------------------------------
# Course Revision — the "Percentage of change in syllabus revision" document
# (BBA Syllabus Revision 2024): its letterhead, orange bands, (A)–(D),
# course-wise by semester, module-wise for every course
# ---------------------------------------------------------------------------

REVISION_BASE = Path(__file__).parent / "doc_templates" / "revision_base.docx"
ORANGE = "FF9933"
YELLOW = "FFFF00"
GREY = "D3D3D3"
THRESHOLD = 20
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


def _rev_font(doc):
    for st in doc.styles:
        if st.name == "Normal":
            st.font.name = "Palatino Linotype"
            st.font.size = Pt(10)


def _orange_band(doc, text, width, size=11, underline=True):
    t = doc.add_table(rows=1, cols=1)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    _write(cell, text, bold=True, size=size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in cell.paragraphs[0].runs:
        r.underline = underline
    _shade(cell, ORANGE)
    _widths(t, [width])
    return t


def _sem_heading(doc, sem):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    label = f"SEMESTER {ROMAN[int(sem) - 1]}" if sem and 0 < int(sem) <= 10 else "SEMESTER NOT GIVEN"
    r = p.add_run(label)
    r.bold = True
    r.underline = True
    r.font.size = Pt(11)


def _pct(v):
    v = _num(v)
    return "—" if v is None else f"{_fmt(v)}%"


def revision_docx(db, dept_code, year, programme_code):
    """The programme's Course Revision in the Office's revision template."""
    from .exporter import revision_summary
    rec = programme_record(db, dept_code, year, programme_code)
    if not rec:
        return None
    dept, prog, cur, _ = rec
    sub = db.submissions.find_one({"dept_code": dept_code, "academic_year": year}) or {}
    stored = ((sub.get("programmes") or {}).get(prog["programme_code"]) or {}).get("prog_revision") or {}
    courses = [c for c in ((stored.get("data") or {}).get("courses") or []) if isinstance(c, dict)]
    name = _title_line(prog, cur)[0]
    rev_year = next((str(c.get("year_latest")) for c in courses if c.get("year_latest")), year[:4])

    doc = Document(str(REVISION_BASE))
    _rev_font(doc)
    sec = doc.sections[0]
    width = int(sec.page_width - sec.left_margin - sec.right_margin)
    s = revision_summary(courses, THRESHOLD)

    # --- page 1: who and when, (A)–(D), the note, the signatures
    _orange_band(doc, "Percentage of change in syllabus revision for all Courses", width * 0.72, underline=False)
    t = doc.add_table(rows=2, cols=4)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, (l1, v1, l2, v2) in enumerate((("NAME OF THE PROGRAMME :", name, "PROGRAMME CODE :", prog["programme_code"]),
                                           ("NAME OF THE DEPARTMENT :", dept.get("dept_name", ""), "YEAR OF REVISION :", rev_year))):
        cells = t.rows[r].cells
        for i, (txt, label) in enumerate(((l1, True), (v1, False), (l2, True), (v2, False))):
            _write(cells[i], txt, bold=True, size=9.5, align=None if label else WD_ALIGN_PARAGRAPH.CENTER)
            if label:
                for run in cells[i].paragraphs[0].runs:
                    run.underline = True
                _shade(cells[i], ORANGE)
    _widths(t, [width * x for x in (0.16, 0.26, 0.15, 0.15)])
    doc.add_paragraph()

    t = doc.add_table(rows=0, cols=3)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, label, v in (("(A)", "Total Number of Courses", str(s["A"])),
                        ("(B)", f"Number of Courses with Syllabus revision above {THRESHOLD}%", str(s["B"])),
                        ("(C)", "Percentage of Courses revised\nFormula: (B / A) x 100", _fmt(s["C"])),
                        ("(D)", "Average Percentage of Syllabus revised considering the percentage of "
                                "syllabus revision in each course *", _pct(s["D"]))):
        cells = t.add_row().cells
        _write(cells[0], k, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[1], label, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(cells[2], v, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        if k == "(D)":
            _shade(cells[2], YELLOW)
    _widths(t, [width * 0.06, width * 0.56, width * 0.1])
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=1)
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _write(t.rows[0].cells[0],
           f"* In {name}, out of {s['A']} number of Courses, {s['B']} number of Courses have undergone syllabus "
           f"revision with more than {THRESHOLD}% is considered and the same is highlighted in YELLOW COLOUR. "
           f"And, the other courses in which the syllabus revision is less than {THRESHOLD}% are highlighted in "
           "GREY COLOUR which are not considered in the above-mentioned percentage of syllabus revision.",
           bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _widths(t, [width * 0.62])
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(40)
    r = p.add_run("Head of the Department" + "\t" * 6 + "Director / Deputy Director of the School")
    r.bold = True

    if not courses:
        _page_break(doc)
        _para(doc, "No Course Revision has been entered for this programme yet.", size=11)
        doc.core_properties.title = _title(f"{name} — Course Revision {rev_year}")
        return _save(doc)

    pairs = list(zip(s["rows"], s["avgs"]))
    sems = sorted({_num(c.get("semester")) for c, _ in pairs}, key=lambda x: (x is None, x or 0))

    # --- course-wise, semester by semester
    _page_break(doc)
    _orange_band(doc, "PERCENTAGE OF CHANGE IN SYLLABUS – COURSE-WISE FOR ALL SEMESTERS", width)
    for sem in sems:
        _sem_heading(doc, sem)
        t = doc.add_table(rows=1, cols=4)
        _borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(("SL", "COURSE CODE", "COURSE TITLE", "PERCENTAGE OF CHANGE IN SYLLABUS")):
            _write(t.rows[0].cells[i], h, bold=True, size=9.5)
            for run in t.rows[0].cells[i].paragraphs[0].runs:
                run.underline = True
            _shade(t.rows[0].cells[i], ORANGE)
        for n, (c, avg) in enumerate([x for x in pairs if _num(x[0].get("semester")) == sem], 1):
            cells = t.add_row().cells
            mark = YELLOW if (avg or 0) > THRESHOLD else GREY
            for i, v in enumerate((str(n), c.get("course_code") or "", str(c.get("course_title") or "").upper(),
                                   _pct(avg))):
                _write(cells[i], v, bold=True, size=9.5)
                _shade(cells[i], mark)
        _widths(t, [width * x for x in (0.06, 0.18, 0.52, 0.18)])

    # --- module-wise, course by course
    _page_break(doc)
    _orange_band(doc, "PERCENTAGE OF CHANGE IN SYLLABUS – MODULE-WISE FOR ALL COURSES", width * 0.72)
    for sem in sems:
        _sem_heading(doc, sem)
        t = doc.add_table(rows=1, cols=2)
        _borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _write(t.rows[0].cells[0], "NAME OF THE PROGRAMME :", bold=True, size=10)
        t.rows[0].cells[0].paragraphs[0].runs[0].underline = True
        _write(t.rows[0].cells[1], name, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _widths(t, [width * 0.2, width * 0.22])
        for c, avg in [x for x in pairs if _num(x[0].get("semester")) == sem]:
            doc.add_paragraph()
            t = doc.add_table(rows=1, cols=4)
            _borders(t)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(("", "YEAR OF PREVIOUS REVISION IN A SUBJECT / COURSE",
                                   "YEAR OF LATEST REVISION OF A SUBJECT / COURSE", "PERCENTAGE OF CHANGE")):
                _write(t.rows[0].cells[i], h, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                for run in t.rows[0].cells[i].paragraphs[0].runs:
                    run.underline = True
                _shade(t.rows[0].cells[i], ORANGE)
            for label, prev, new in (("YEAR", c.get("year_previous") or "—", c.get("year_latest") or rev_year),
                                     ("SUBJECT / COURSE TITLE", (c.get("prev_title") or "—").upper(),
                                      str(c.get("course_title") or "").upper()),
                                     ("SUBJECT CODE", c.get("prev_code") or "—", c.get("course_code") or "")):
                cells = t.add_row().cells
                _write(cells[0], label, bold=True, size=9.5)
                cells[0].paragraphs[0].runs[0].underline = True
                _write(cells[1], prev, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
                _write(cells[2], new, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            for n, m in enumerate([m for m in c.get("modules") or [] if isinstance(m, dict)], 1):
                cells = t.add_row().cells
                pct = _num(m.get("pct"))
                _write(cells[0], f"Module {ROMAN[n - 1] if n <= 10 else n}", size=9.5,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
                _write(cells[1], m.get("previous") or "", size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                _write(cells[2], m.get("revised") or "", size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                _write(cells[3], _pct(pct), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
                if str(m.get("revised") or "").strip() and (pct or 0) > 0:
                    _shade(cells[2], YELLOW)
            cells = t.add_row().cells
            _write(cells[0], "AVERAGE PERCENTAGE ON REVISION CONSIDERING ALL MODULES", bold=True, size=9.5)
            _write(cells[3], _pct(avg), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade(cells[3], YELLOW if (avg or 0) > THRESHOLD else GREY)
            _widths(t, [width * x for x in (0.16, 0.34, 0.34, 0.12)])

    doc.core_properties.title = _title(f"{name} — Course Revision {rev_year}")
    return _save(doc)


# ---------------------------------------------------------------------------
# the same document as a web page (the share viewer)
# ---------------------------------------------------------------------------

def _cell_view(tc, doc):
    from docx.table import _Cell
    cell = _Cell(tc, doc)
    pr = tc.tcPr
    span, fill = 1, None
    if pr is not None:
        gs = pr.find(qn("w:gridSpan"))
        span = int(gs.get(qn("w:val"))) if gs is not None else 1
        shd = pr.find(qn("w:shd"))
        fill = shd.get(qn("w:fill")) if shd is not None else None
    lines = []
    for p in cell.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs and not p.text:
            continue
        lines.append({"text": p.text, "bold": bool(runs) and all(r.bold for r in runs),
                      "color": next((str(r.font.color.rgb) for r in runs
                                     if r.font.color is not None and r.font.color.type is not None), None),
                      "align": {1: "center", 2: "right", 3: "justify"}.get(
                          int(p.alignment) if p.alignment is not None else 0)})
    return {"span": span, "fill": fill if fill and fill.lower() not in ("auto", "ffffff") else None,
            "lines": lines}


def to_view(buf):
    """A generated document as blocks a page can show in order:
    ("p", {text, bold, size, color, align}), ("break", None) and
    ("table", [[cell, …], …]) with each cell's span, fill and lines."""
    from docx.text.paragraph import Paragraph
    doc = Document(buf)
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            if child.xpath(".//w:br[@w:type='page']"):
                out.append(("break", None))
                continue
            p = Paragraph(child, doc)
            runs = [r for r in p.runs if r.text]
            if not runs:
                continue
            r = runs[0]
            out.append(("p", {"text": p.text, "bold": bool(r.bold),
                              "size": r.font.size.pt if r.font.size else 10,
                              "color": str(r.font.color.rgb) if r.font.color is not None
                              and r.font.color.type is not None else None,
                              "align": "center" if p.alignment == WD_ALIGN_PARAGRAPH.CENTER else "left"}))
        elif tag == "tbl":
            rows = []
            for tr in child.iterchildren(qn("w:tr")):
                rows.append([_cell_view(tc, doc) for tc in tr.iterchildren(qn("w:tc"))])
            out.append(("table", rows))
    header = " ".join(p.text for p in doc.sections[0].header.paragraphs if p.text).strip()
    return {"header": header, "blocks": out,
            "landscape": doc.sections[0].page_width > doc.sections[0].page_height}
