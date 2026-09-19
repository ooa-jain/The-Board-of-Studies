"""Excel and Word exports of the stored submissions."""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ugc_rules as U
from .db import get_db
from .schema import STAGE_BY_KEY, STAGES
from .workflow import compute_status, progress

NAVY = "0F2A4A"
GOLD = "C8A44B"
LIGHT = "F3F6FA"

_thin = Side(style="thin", color="D5DEE8")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _head(ws, row, values, fill=NAVY, color="FFFFFF"):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color=color, size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 26


def _autosize(ws, max_width=52):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(12, width + 3), max_width)


def _flat(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple)):
        return "; ".join(_flat(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_flat(v)}" for k, v in value.items())
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")
    return str(value)


# ---------------------------------------------------------------------------
# Excel — one department
# ---------------------------------------------------------------------------

def department_excel(dept_code: str, year: str) -> io.BytesIO:
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or {}
    sub = db.submissions.find_one({"dept_code": dept_code, "academic_year": year}) or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "JAIN (Deemed-to-be University) — Office of Academics"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = f"Board of Studies Data Repository · {year}"
    ws["A2"].font = Font(size=11, color="53627A")

    meta = [
        ("Department", dept.get("dept_name")),
        ("Department code", dept.get("dept_code")),
        ("School / Faculty", dept.get("school")),
        ("Campus", dept.get("campus")),
        ("Overall progress", f"{progress(sub)['done']} of {progress(sub)['total']} stages"),
        ("Generated on", datetime.now().strftime("%d %b %Y, %H:%M")),
    ]
    r = 4
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        ws.cell(row=r, column=2, value=_flat(v))
        r += 1

    r += 1
    _head(ws, r, ["Stage", "Group", "Status", "Errors", "Warnings", "Submitted on"])
    r += 1
    for s in STAGES:
        st = compute_status(sub, s["key"])
        state = (sub.get("stages") or {}).get(s["key"], {})
        summary = state.get("summary") or {}
        for i, v in enumerate([s["title"], s["group"], st.title(),
                               summary.get("errors", 0), summary.get("warnings", 0),
                               _flat(state.get("submitted_at"))], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i == 3:
                c.fill = PatternFill("solid", fgColor={"submitted": "DFF3E4",
                                                       "returned": "FDE7E7",
                                                       "locked": "EFEFEF"}.get(st, "FFF7E0"))
        r += 1
    _autosize(ws)

    # one sheet per stage
    for s in STAGES:
        state = (sub.get("stages") or {}).get(s["key"], {})
        data = state.get("data") or {}
        if not data:
            continue
        title = s["title"][:28]
        sheet = wb.create_sheet(title)
        row = 1
        sheet.cell(row=row, column=1, value=s["title"]).font = Font(bold=True, size=13, color=NAVY)
        row += 2
        for section in s["sections"]:
            sk = section["key"]
            payload = data.get(sk)
            if not payload:
                continue
            sheet.cell(row=row, column=1, value=section["title"]).font = Font(bold=True, size=11)
            row += 1
            if section.get("type") == "table" and isinstance(payload, list):
                cols = section.get("columns", [])
                _head(sheet, row, [c["label"] for c in cols], fill=GOLD, color="1A1A1A")
                row += 1
                for item in payload:
                    for i, c in enumerate(cols, start=1):
                        cell = sheet.cell(row=row, column=i, value=_flat(item.get(c["name"])))
                        cell.border = BORDER
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    row += 1
            elif isinstance(payload, dict):
                for k, v in payload.items():
                    sheet.cell(row=row, column=1, value=k).font = Font(bold=True, size=10)
                    sheet.cell(row=row, column=2, value=_flat(v)).alignment = \
                        Alignment(wrap_text=True, vertical="top")
                    row += 1
            row += 1
        _autosize(sheet)

    # per-programme sheets
    for code, stages in (sub.get("programmes") or {}).items():
        for skey, state in stages.items():
            data = state.get("data") or {}
            if not data:
                continue
            sdef = STAGE_BY_KEY.get(skey, {})
            sheet = wb.create_sheet(f"{code}-{skey}"[:31])
            row = 1
            sheet.cell(row=row, column=1,
                       value=f"{code} — {sdef.get('title', skey)}").font = \
                Font(bold=True, size=13, color=NAVY)
            row += 2
            for section in sdef.get("sections", []):
                payload = data.get(section["key"])
                if not payload:
                    continue
                sheet.cell(row=row, column=1, value=section["title"]).font = Font(bold=True, size=11)
                row += 1
                if section.get("type") == "table" and isinstance(payload, list):
                    cols = section.get("columns", [])
                    _head(sheet, row, [c["label"] for c in cols], fill=GOLD, color="1A1A1A")
                    row += 1
                    for item in payload:
                        for i, c in enumerate(cols, start=1):
                            cell = sheet.cell(row=row, column=i, value=_flat(item.get(c["name"])))
                            cell.border = BORDER
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                        row += 1
                elif isinstance(payload, dict):
                    for k, v in payload.items():
                        label = next((r["label"] for r in U.DEFAULT_TABLE_2 if r["key"] == k), k)
                        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
                        sheet.cell(row=row, column=2, value=_flat(v))
                        row += 1
                row += 1
            _autosize(sheet)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Excel — the whole institution
# ---------------------------------------------------------------------------

def institution_excel(year: str) -> io.BytesIO:
    db = get_db()
    depts = list(db.departments.find({"active": True}).sort([("campus", 1), ("dept_name", 1)]))
    subs = {s["dept_code"]: s for s in db.submissions.find({"academic_year": year})}

    wb = Workbook()
    ws = wb.active
    ws.title = "Status"
    ws["A1"] = f"BoS Data Repository — institution status, {year}"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)

    header = ["Campus", "School", "Department", "Code", "Progress"] + \
             [s["title"] for s in STAGES]
    _head(ws, 3, header)
    r = 4
    for d in depts:
        sub = subs.get(d["dept_code"], {})
        p = progress(sub) if sub else {"done": 0, "total": len(STAGES)}
        row = [d.get("campus"), d.get("school"), d.get("dept_name"), d.get("dept_code"),
               f"{p['done']}/{p['total']}"]
        for s in STAGES:
            row.append(compute_status(sub, s["key"]).title() if sub else "Not started")
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=_flat(v))
            c.border = BORDER
            if i > 6:
                c.fill = PatternFill("solid", fgColor={"Submitted": "DFF3E4",
                                                       "Returned": "FDE7E7",
                                                       "Draft": "FFF7E0",
                                                       "Locked": "EFEFEF"}.get(v, "FFFFFF"))
                c.alignment = Alignment(horizontal="center")
        r += 1
    ws.freeze_panes = "G4"
    _autosize(ws, max_width=28)

    # credit compliance across programmes
    cs = wb.create_sheet("Credit compliance")
    _head(cs, 1, ["Campus", "Department", "Programme", "Code", "Track", "Declared total",
                  "Required", "Shortfall", "Errors"])
    r = 2
    rules = db.rules.find_one({"_id": "ugc"}) or {}
    totals = rules.get("totals") or U.DEFAULT_TOTALS
    for d in depts:
        sub = subs.get(d["dept_code"])
        if not sub:
            continue
        progs = ((sub.get("stages") or {}).get("ugc_programme", {})
                 .get("data", {}).get("programmes") or [])
        for p in progs:
            pcode = p.get("programme_code")
            state = ((sub.get("programmes") or {}).get(pcode, {})
                     .get("ugc_curriculum", {}))
            matrix = (state.get("data") or {}).get("credit_summary") or {}
            track = U.get_track(p.get("degree_level"))
            declared = sum(float(v) for k, v in matrix.items()
                           if k != "total" and str(v).replace(".", "", 1).isdigit())
            required = totals.get(track) if track else None
            shortfall = (required - declared) if required else ""
            errors = (state.get("summary") or {}).get("errors", "")
            for i, v in enumerate([d.get("campus"), d.get("dept_name"), p.get("programme_name"),
                                   pcode, U.TRACKS.get(track, "—"), declared or "",
                                   required or "—", shortfall if shortfall and shortfall > 0 else "",
                                   errors], start=1):
                c = cs.cell(row=r, column=i, value=_flat(v))
                c.border = BORDER
            r += 1
    _autosize(cs)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Word report
# ---------------------------------------------------------------------------

def submission_word(dept_code: str, year: str) -> io.BytesIO:
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or {}
    sub = db.submissions.find_one({"dept_code": dept_code, "academic_year": year}) or {}

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("JAIN (Deemed-to-be University)")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)

    sub_t = doc.add_paragraph()
    sub_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub_t.add_run("Office of Academics · Board of Studies Data Repository")
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0x53, 0x62, 0x7A)

    doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in [("Department", dept.get("dept_name")),
                 ("Department code", dept.get("dept_code")),
                 ("School / Faculty", dept.get("school")),
                 ("Campus", dept.get("campus")),
                 ("Academic year", year),
                 ("Report generated", datetime.now().strftime("%d %B %Y, %H:%M"))]:
        row = t.add_row().cells
        row[0].text = k
        row[1].text = _flat(v)

    for s in STAGES:
        state = (sub.get("stages") or {}).get(s["key"], {})
        data = state.get("data") or {}
        if not data:
            continue
        doc.add_heading(s["title"], level=1)
        status = compute_status(sub, s["key"])
        p = doc.add_paragraph()
        pr = p.add_run(f"Status: {status.title()}")
        pr.italic = True
        pr.font.size = Pt(9)

        for section in s["sections"]:
            payload = data.get(section["key"])
            if not payload:
                continue
            doc.add_heading(section["title"], level=2)
            if section.get("type") == "table" and isinstance(payload, list):
                cols = section.get("columns", [])
                keep = [c for c in cols if c.get("type") != "file"][:8]
                tbl = doc.add_table(rows=1, cols=len(keep))
                tbl.style = "Light Grid Accent 1"
                for i, c in enumerate(keep):
                    cell = tbl.rows[0].cells[i]
                    cell.text = c["label"]
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                for item in payload:
                    cells = tbl.add_row().cells
                    for i, c in enumerate(keep):
                        cells[i].text = _flat(item.get(c["name"]))[:300]
            elif isinstance(payload, dict):
                for k, v in payload.items():
                    para = doc.add_paragraph()
                    label = next((r["label"] for r in U.DEFAULT_TABLE_2 if r["key"] == k), k)
                    para.add_run(f"{label}: ").bold = True
                    para.add_run(_flat(v))

    for code, stages in (sub.get("programmes") or {}).items():
        doc.add_page_break()
        doc.add_heading(f"Programme {code}", level=1)
        for skey, state in stages.items():
            data = state.get("data") or {}
            if not data:
                continue
            sdef = STAGE_BY_KEY.get(skey, {})
            doc.add_heading(sdef.get("title", skey), level=2)
            for section in sdef.get("sections", []):
                payload = data.get(section["key"])
                if not payload:
                    continue
                doc.add_heading(section["title"], level=3)
                if isinstance(payload, list):
                    cols = [c for c in section.get("columns", []) if c.get("type") != "file"][:8]
                    tbl = doc.add_table(rows=1, cols=len(cols))
                    tbl.style = "Light Grid Accent 1"
                    for i, c in enumerate(cols):
                        tbl.rows[0].cells[i].text = c["label"]
                    for item in payload:
                        cells = tbl.add_row().cells
                        for i, c in enumerate(cols):
                            cells[i].text = _flat(item.get(c["name"]))[:300]
                elif isinstance(payload, dict):
                    for k, v in payload.items():
                        label = next((r["label"] for r in U.DEFAULT_TABLE_2 if r["key"] == k), k)
                        para = doc.add_paragraph()
                        para.add_run(f"{label}: ").bold = True
                        para.add_run(_flat(v))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
