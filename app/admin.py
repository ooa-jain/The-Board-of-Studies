"""Office of Academic Affairs administration console."""

from __future__ import annotations

import io
import re
from datetime import datetime

from flask import (Blueprint, abort, current_app, flash, jsonify, redirect,
                   render_template, request, send_file, session, url_for)

from . import ugc_rules as U
from .auth import admin_required
from .db import audit, get_db, issue_department_login, now, rules_doc
from .exporter import (department_excel, institution_excel, submission_word)
from .importer import parse_workbook
from .packs import available_packs, load_pack
from .report import AdminLinks, ShareLinks
from .schema import STAGE_BY_KEY, STAGES
from .workflow import (compute_status, department_analysis, get_or_create_submission,
                       institution_analysis, programmes_of, progress,
                       stage_analysis, return_stage, stage_board, unlock_stage, batches)

bp = Blueprint("admin", __name__)


def _year():
    from .db import settings
    return settings().get("academic_year") or current_app.config["ACADEMIC_YEAR"]


def _actor():
    return session["user"]["username"]


# ---------------------------------------------------------------------------
# dashboard
# ---------------------------------------------------------------------------

def _activity(db, days_back=14):
    """Updates from departments, a day at a time, for the last fortnight."""
    from datetime import timedelta
    today = now().replace(hour=0, minute=0, second=0, microsecond=0)
    days = [today - timedelta(days=i) for i in range(days_back - 1, -1, -1)]
    per_day = {d.date(): {"all": 0, "submitted": 0} for d in days}
    for n in db.notifications.find({"at": {"$gte": days[0]}}, {"at": 1, "event": 1}):
        k = n["at"].date()
        if k in per_day:
            per_day[k]["all"] += 1
            if n.get("event") == "submitted":
                per_day[k]["submitted"] += 1
    series = [{"label": d.strftime("%a %d %b"), "day": d.strftime("%a")[:2], "date": d.strftime("%d"),
               **per_day[d.date()]} for d in days]
    peak = max([x["all"] for x in series] + [0])
    return {"series": series, "peak": peak, "total": sum(x["all"] for x in series),
            "submitted": sum(x["submitted"] for x in series)}


def _picture(db, year):
    """The whole institution in one reading — for the home and for the
    one-click analysis, so the two can never disagree."""
    departments = list(db.departments.find({"active": True})
                       .sort([("place", 1), ("campus", 1), ("dept_name", 1)]))
    subs = {s["dept_code"]: s for s in db.submissions.find({"academic_year": year})}
    users = {u.get("dept_code"): u for u in db.users.find({"role": "department"})}
    from .workflow import programme_progress
    rows = []
    for d in departments:
        sub = subs.get(d["dept_code"])
        r = department_analysis(d, sub, users.get(d["dept_code"]))
        r["programmes"] = programme_progress(sub or {}, d)
        r["prog_done"] = sum(1 for p in r["programmes"] if p["complete"])
        rows.append(r)
    totals = institution_analysis(rows)
    stages = stage_analysis([r["board"] for r in rows])
    for st in stages:
        st["in_progress"] = st.get("draft", 0)
        st["not_started"] = st.get("open", 0) + st.get("locked", 0)

    progs = [p for r in rows for p in r["programmes"]]
    levels = []
    for lv in ("UG", "PG"):
        mine = [p for p in progs if (p["level"] or "").upper().startswith(lv)]
        if mine:
            levels.append({"level": lv, "total": len(mine),
                           "complete": sum(1 for p in mine if p["complete"]),
                           "started": sum(1 for p in mine if p["started"])})
    revs = [p["revision"] for p in progs if p["revision"] is not None]
    programmes = {"total": len(progs), "complete": sum(1 for p in progs if p["complete"]),
                  "started": sum(1 for p in progs if p["started"]), "levels": levels,
                  "revision": round(sum(revs) / len(revs), 1) if revs else None,
                  "revised": len(revs)}

    def group(key):
        out = {}
        for r in rows:
            g = out.setdefault(r["dept"].get(key) or "—", {"name": r["dept"].get(key) or "—", "departments": 0,
                                                          "complete": 0, "done": 0, "total": 0})
            g["departments"] += 1
            g["complete"] += 1 if r["state"] == "complete" else 0
            g["done"] += r["done"]
            g["total"] += r["total"]
        for g in out.values():
            g["percent"] = round(g["done"] * 100 / g["total"]) if g["total"] else 0
        return sorted(out.values(), key=lambda g: (-g["percent"], g["name"]))

    attention = {"comments": db.comments.count_documents({"academic_year": year, "status": "open"}),
                 "returned": sum(r["returned"] for r in rows),
                 "bad_files": db.files.count_documents({"academic_year": year, "$or": [
                     {"keyword_match.status": "miss"}, {"keyword_match.looks_like": {"$exists": True}}]})}
    attention["total"] = attention["comments"] + attention["returned"] + attention["bad_files"]
    files = db.files.count_documents({"academic_year": year})
    return {"rows": rows, "totals": totals, "stages": stages, "programmes": programmes,
            "by_campus": group("campus"), "by_school": group("school"), "attention": attention,
            "files": files, "activity": _activity(db)}


_STATE_RANK = {"complete": 5, "in_progress": 4, "returned": 3, "not_started": 2, "never_in": 1, "no_login": 0}


def _ahead(r):
    """How far a department has got: stages completed, then programmes
    completed, then work begun (half-filled stages, programmes started)."""
    return (r["percent"], r["prog_done"], r["half"] + r["returned"],
            sum(1 for p in r["programmes"] if p["started"]), _STATE_RANK.get(r["state"], 0))


def _findings(pic):
    """What the numbers say, in sentences, most pressing first."""
    t, rows, out = pic["totals"], pic["rows"], []
    n = t["departments"]
    if not n:
        return ["No departments yet — import them to begin."]
    out.append(f"{t['percent']}% of all stages are completed — {t['stages_done']} of {t['stages_total']} "
               f"across {n} department{'s' if n != 1 else ''}.")
    if t["complete"]:
        out.append(f"{t['complete']} of {n} departments have completed every stage.")
    idle = [r["dept"]["dept_name"] for r in rows if r["state"] in ("not_started", "never_in", "no_login")]
    if idle:
        names = ", ".join(idle[:4]) + (f" and {len(idle) - 4} more" if len(idle) > 4 else "")
        out.append(f"{len(idle)} department{'s have' if len(idle) != 1 else ' has'} filed nothing yet: {names}.")
    if t["returned"]:
        out.append(f"{t['returned']} department{'s have' if t['returned'] != 1 else ' has'} a stage sent back for correction.")
    stages = pic["stages"]
    if stages:
        slow = min(stages, key=lambda s: (s["submitted"], -s["n"]))
        fast = max(stages, key=lambda s: (s["submitted"], -s["n"]))
        if fast["submitted"] != slow["submitted"]:
            out.append(f"Slowest stage: {slow['title']} — {slow['submitted']} of {n} completed. "
                       f"Furthest along: {fast['title']} — {fast['submitted']} of {n}.")
        elif not fast["submitted"]:
            busy = max(s["in_progress"] for s in stages)
            out.append("No stage has been completed by any department yet"
                       + (f"; {busy} department{'s are' if busy != 1 else ' is'} filling stages now." if busy else "."))
    lead = max(rows, key=_ahead)
    if lead["state"] not in ("not_started", "never_in", "no_login"):
        progs = len(lead["programmes"])
        started = sum(1 for p in lead["programmes"] if p["started"])
        out.append(f"Furthest ahead: {lead['dept']['dept_name']} — {lead['done']} of {lead['total']} stages completed"
                   + (f", {lead['prog_done']} of {progs} programmes completed ({started} started)." if progs else "."))
    pr = pic["programmes"]
    if pr["total"]:
        lv = " · ".join(f"{x['level']} {x['complete']} of {x['total']}" for x in pr["levels"])
        out.append(f"Programmes completed: {pr['complete']} of {pr['total']}" + (f" ({lv})." if lv else "."))
    if pr["revision"] is not None:
        out.append(f"Average syllabus change reported in Course Revision: {pr['revision']}% "
                   f"across {pr['revised']} programme{'s' if pr['revised'] != 1 else ''}.")
    camp = pic["by_campus"]
    if len(camp) > 1 and camp[0]["percent"] != camp[-1]["percent"]:
        out.append(f"By campus, {camp[0]['name']} leads at {camp[0]['percent']}%; "
                   f"{camp[-1]['name']} trails at {camp[-1]['percent']}%.")
    a = pic["attention"]
    if a["total"]:
        out.append(f"Needs attention: {a['bad_files']} file{'s' if a['bad_files'] != 1 else ''} to check, "
                   f"{a['returned']} sent back, {a['comments']} open comment{'s' if a['comments'] != 1 else ''}.")
    return out


@bp.route("/")
@admin_required
def dashboard():
    db = get_db()
    year = _year()
    pic = _picture(db, year)
    # the departments furthest behind come first: they are the ones to chase
    rows = sorted(pic["rows"], key=lambda r: (_ahead(r), r["dept"].get("dept_name", "")))
    return render_template("admin/dashboard.html", year=year, pic=pic, rows=rows,
                           latest=list(db.notifications.find().sort("at", -1).limit(8)))


@bp.route("/analysis/report")
@admin_required
def analysis_report():
    """Analysis in one click: everything on one page, in charts and in
    sentences, ready to print or save as PDF."""
    db = get_db()
    year = _year()
    pic = _picture(db, year)
    ranked = sorted(pic["rows"], key=lambda r: tuple(-x for x in _ahead(r)) + (r["dept"].get("dept_name", ""),))
    return render_template("admin/analysis_report.html", year=year, pic=pic, ranked=ranked,
                           findings=_findings(pic), made=now())


# ---------------------------------------------------------------------------
# analysis
# ---------------------------------------------------------------------------

@bp.route("/flow")
@admin_required
def flow():
    """How one department's data moves from stage to stage, in 3D. The
    department — and a programme of it — is chosen in the side panel."""
    departments = list(get_db().departments.find({"active": True}).sort("dept_name", 1))
    return render_template("admin/flow.html", year=_year(), departments=departments)


def _flow_for(dept_code, programme_code=None):
    from .flows import flows_for
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    sub = db.submissions.find_one({"dept_code": dept_code, "academic_year": _year()}) or {}
    return dept, sub, flows_for(sub, dept, programme_code)


@bp.route("/flow.json")
@admin_required
def flow_data():
    code = request.args.get("dept") or abort(400)
    _, _, data = _flow_for(code, request.args.get("prog") or None)
    return jsonify(data)


@bp.route("/flow.xlsx")
@admin_required
def flow_excel():
    """One department's data mapping — every flow, for every programme."""
    from .exporter import flow_mapping_excel
    from .flows import flows_for
    code = request.args.get("dept") or abort(400)
    dept, sub, data = _flow_for(code)
    per_prog = [(p, flows_for(sub, dept, p["code"])["flows"]) for p in data["programmes"]]
    buf = flow_mapping_excel(data, per_prog, _year())
    return send_file(buf, as_attachment=True,
                     download_name=f"BoS-data-mapping-{code}-{_year()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/analysis")
@admin_required
def analysis():
    """Every department in one reading: done, part-done, untouched, failing.

    The Office's standing question is "who still owes me what", and it used to
    be answered by opening departments one at a time. This answers it once.
    """
    db = get_db()
    year = _year()

    # Demonstration mode. Everything it needs lives in app/demo.py; delete that
    # module and this branch and the page is live-only.
    if request.args.get("demo") == "1":
        from .demo import demo_analysis
        rows, totals, stages = demo_analysis()
        return render_template("admin/analysis.html", rows=rows, totals=totals,
                               stages=stages, year=year, demo=True,
                               filters={"state": None, "q": ""})

    departments = list(db.departments.find({"active": True})
                       .sort([("place", 1), ("campus", 1), ("dept_name", 1)]))
    subs = {s["dept_code"]: s for s in db.submissions.find({"academic_year": year})}
    users = {u.get("dept_code"): u for u in db.users.find({"role": "department"})}

    rows = [department_analysis(d, subs.get(d["dept_code"]), users.get(d["dept_code"]))
            for d in departments]
    totals = institution_analysis(rows)
    stages = stage_analysis([stage_board(subs.get(d["dept_code"]) or {})
                             for d in departments])

    # filtering happens after the totals, so the summary always describes the
    # whole institution and not whatever slice is on screen
    state = request.args.get("state")
    search = (request.args.get("q") or "").strip().lower()
    if state:
        rows = [r for r in rows if r["state"] == state]
    if search:
        rows = [r for r in rows
                if search in (r["dept"].get("dept_name", "") + " "
                              + r["dept"].get("school", "") + " "
                              + r["dept"].get("dept_code", "")).lower()]

    return render_template("admin/analysis.html", rows=rows, totals=totals,
                           stages=stages, year=year, demo=False,
                           filters={"state": state, "q": search})


# ---------------------------------------------------------------------------
# department master
# ---------------------------------------------------------------------------

@bp.route("/departments")
@admin_required
def departments():
    db = get_db()
    q = {}
    place = request.args.get("place")
    campus = request.args.get("campus")
    school = request.args.get("school")
    search = (request.args.get("q") or "").strip()
    if place:
        q["place"] = place
    if campus:
        q["campus"] = campus
    if school:
        q["school"] = school
    if search:
        q["$or"] = [{"dept_name": {"$regex": search, "$options": "i"}},
                    {"dept_code": {"$regex": search, "$options": "i"}},
                    {"school": {"$regex": search, "$options": "i"}},
                    {"faculty": {"$regex": search, "$options": "i"}}]
    depts = list(db.departments.find(q)
                 .sort([("place", 1), ("campus", 1), ("school", 1), ("dept_name", 1)]))

    # only the campuses in the city being looked at: offering Kochi Campus
    # while Bangalore is selected would be a filter that can only return none
    campuses = [c for c in current_app.config["CAMPUSES"]
                if not place or db.departments.count_documents(
                    {"place": place, "campus": c})]
    return render_template("admin/departments.html", departments=depts,
                           demo=db.departments.find_one({"dept_code": "DEMO"}),
                           packs=available_packs(),
                           schools=sorted(x for x in db.departments.distinct("school") if x),
                           campuses=campuses,
                           places=current_app.config["PLACES"],
                           total=db.departments.count_documents({}),
                           filters={"place": place, "campus": campus,
                                    "school": school, "q": search})


@bp.route("/departments/new", methods=["GET", "POST"])
@bp.route("/departments/<dept_code>/edit", methods=["GET", "POST"])
@admin_required
def department_form(dept_code=None):
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) if dept_code else None
    if dept_code and not dept:
        abort(404)

    # existing values offered as suggestions, so a new department lands in an
    # existing faculty and school rather than a near-miss spelling of one
    def _form(d):
        return render_template(
            "admin/department_form.html", dept=d,
            campuses=current_app.config["CAMPUSES"],
            places=current_app.config["PLACES"],
            faculties=sorted(x for x in db.departments.distinct("faculty") if x),
            schools=sorted(x for x in db.departments.distinct("school") if x))

    if request.method == "POST":
        f = request.form
        code = (f.get("dept_code") or "").strip().upper()
        if not code:
            flash("Department code is required.", "error")
            return _form(dept or f)
        clash = db.departments.find_one({"dept_code": code})
        if clash and (not dept or clash["_id"] != dept["_id"]):
            flash(f"Department code “{code}” is already in use.", "error")
            return _form(f)

        payload = {
            "dept_code": code,
            "dept_name": (f.get("dept_name") or "").strip(),
            "faculty": (f.get("faculty") or "").strip(),
            "school": (f.get("school") or "").strip(),
            "campus": f.get("campus") or current_app.config["CAMPUSES"][0],
            "place": f.get("place") or current_app.config["PLACES"][0],
            "active": f.get("active") == "on",
            "updated_at": now(),
        }
        if dept:
            db.departments.update_one({"_id": dept["_id"]}, {"$set": payload})
            audit(_actor(), "department.updated", code)
            flash(f"{payload['dept_name']} has been updated.", "success")
        else:
            payload["created_at"] = now()
            db.departments.insert_one(payload)
            audit(_actor(), "department.created", code)
            flash(f"{payload['dept_name']} has been added. "
                  f"Use “Issue every missing login” on the department list "
                  f"to give it one.", "success")
        return redirect(url_for("admin.departments"))

    return _form(dept)


@bp.route("/departments/<dept_code>/toggle", methods=["POST"])
@admin_required
def department_toggle(dept_code):
    db = get_db()
    d = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    new = not d.get("active", True)
    db.departments.update_one({"_id": d["_id"]}, {"$set": {"active": new}})
    db.users.update_one({"dept_code": dept_code}, {"$set": {"active": new}})
    audit(_actor(), "department.active" if new else "department.disabled", dept_code)
    flash(f"{d['dept_name']} is now {'active' if new else 'disabled'}.", "info")
    return redirect(request.referrer or url_for("admin.departments"))


# ---------------------------------------------------------------------------
# credentials
# ---------------------------------------------------------------------------

@bp.route("/departments/<dept_code>/credentials", methods=["POST"])
@admin_required
def generate_credentials(dept_code):
    """Create or reset the department login. Username is derived from the code."""
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)

    had_login = bool(dept.get("username"))
    username, password = issue_department_login(db, dept, actor=_actor())

    audit(_actor(), "credentials.reset" if had_login else "credentials.generated", dept_code)

    if request.headers.get("Accept", "").startswith("application/json"):
        return jsonify({"ok": True, "username": username, "password": password})
    flash(f"Login generated for {dept['dept_name']}.", "success")
    return redirect(url_for("admin.credential_slip", dept_code=dept_code))


@bp.route("/departments/<dept_code>/slip")
@admin_required
def credential_slip(dept_code):
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    if not dept.get("initial_password"):
        flash("This department has already signed in, so the password is no longer visible. "
              "Reset it if they need a new one.", "info")
        return redirect(url_for("admin.departments"))
    return render_template("admin/credential_slip.html", dept=dept,
                           login_url=url_for("auth.login", _external=True), year=_year())


@bp.route("/departments/clear", methods=["POST"])
@admin_required
def departments_clear():
    """Remove every department in one go, with its login and its submissions.

    This empties the master, so it asks for the word to be typed rather than
    for a button to be clicked. It is not as final as it sounds: seed.py puts
    the whole list back from the Office of Academic Affairs workbook.
    """
    db = get_db()
    if (request.form.get("confirm") or "").strip().upper() != "REMOVE ALL":
        flash("Nothing was removed — type REMOVE ALL to confirm.", "warning")
        return redirect(url_for("admin.departments"))

    counts = {
        "departments": db.departments.count_documents({}),
        "logins": db.users.count_documents({"role": "department"}),
        "submissions": db.submissions.count_documents({}),
    }
    db.departments.delete_many({})
    db.users.delete_many({"role": "department"})
    db.submissions.delete_many({})
    audit(_actor(), "departments.cleared", detail=counts)
    flash(f"Removed {counts['departments']} department(s), "
          f"{counts['logins']} login(s) and {counts['submissions']} submission(s). "
          f"Run seed.py to put the master back from the workbook.", "success")
    return redirect(url_for("admin.departments"))


@bp.route("/credentials/bulk", methods=["POST"])
@admin_required
def bulk_credentials():
    """Generate logins for every active department that has none."""
    db = get_db()
    made = 0
    for dept in list(db.departments.find({"active": True, "username": {"$exists": False}})):
        issue_department_login(db, dept, actor=_actor())
        made += 1
    audit(_actor(), "credentials.bulk", detail={"count": made})
    if made:
        flash(f"Issued a login for {made} department(s). The passwords are on "
              f"the credential sheet until each department signs in.", "success")
    else:
        flash("Every active department already has a login.", "info")
    return redirect(url_for("admin.departments"))


@bp.route("/credentials/sheet")
@admin_required
def credential_sheet():
    """Printable sheet of every credential still in the clear."""
    db = get_db()
    depts = list(db.departments.find({"initial_password": {"$exists": True}})
                 .sort([("campus", 1), ("dept_name", 1)]))
    return render_template("admin/credential_sheet.html", departments=depts,
                           login_url=url_for("auth.login", _external=True), year=_year())


# ---------------------------------------------------------------------------
# Excel import
# ---------------------------------------------------------------------------

@bp.route("/import", methods=["GET", "POST"])
@admin_required
def import_departments():
    if request.method == "GET":
        return render_template("admin/import.html", preview=None, meta=None)

    upload = request.files.get("workbook")
    if not upload or not upload.filename:
        flash("Choose a workbook to upload.", "error")
        return redirect(url_for("admin.import_departments"))

    data = upload.read()
    sheet = request.form.get("sheet") or None
    default_campus = (request.form.get("default_campus")
                      or current_app.config["CAMPUSES"][0])
    default_place = (request.form.get("default_place")
                     or current_app.config["PLACES"][0])
    try:
        rows, meta = parse_workbook(data, sheet, default_campus=default_campus,
                                    default_place=default_place)
    except Exception as exc:
        flash(f"That file could not be read: {exc}", "error")
        return redirect(url_for("admin.import_departments"))

    db = get_db()
    existing = {d["dept_code"] for d in db.departments.find({}, {"dept_code": 1})}
    for r in rows:
        r["exists"] = r["dept_code"] in existing

    session["import_rows"] = rows[:500]
    return render_template("admin/import.html", preview=rows, meta=meta,
                           default_campus=default_campus,
                           default_place=default_place,
                           campuses=current_app.config["CAMPUSES"],
                           places=current_app.config["PLACES"])


@bp.route("/import/commit", methods=["POST"])
@admin_required
def import_commit():
    rows = session.pop("import_rows", None)
    if not rows:
        flash("The import preview expired. Upload the workbook again.", "error")
        return redirect(url_for("admin.import_departments"))

    chosen = set(request.form.getlist("include"))
    overwrite = request.form.get("overwrite") == "on"
    make_logins = request.form.get("make_logins") == "on"

    db = get_db()
    added = updated = skipped = logins = 0
    for r in rows:
        if r["dept_code"] not in chosen:
            continue
        existing = db.departments.find_one({"dept_code": r["dept_code"]})
        payload = {k: r[k] for k in
                   ("dept_name", "school", "dept_code", "campus")}
        if r.get("place"):
            payload["place"] = r["place"]
        if r.get("faculty"):
            payload["faculty"] = r["faculty"]
        payload["active"] = True
        payload["updated_at"] = now()
        if existing and not overwrite:
            skipped += 1
            continue
        if existing:
            db.departments.update_one({"_id": existing["_id"]}, {"$set": payload})
            updated += 1
        else:
            payload["created_at"] = now()
            db.departments.insert_one(payload)
            added += 1

        if make_logins:
            dept = db.departments.find_one({"dept_code": r["dept_code"]})
            if not dept.get("username"):
                issue_department_login(db, dept, actor=_actor())
                logins += 1

    audit(_actor(), "departments.imported",
          detail={"added": added, "updated": updated, "skipped": skipped, "logins": logins})
    flash(f"Import finished — {added} added, {updated} updated, {skipped} skipped, "
          f"{logins} login(s) generated.", "success")
    return redirect(url_for("admin.departments"))


# ---------------------------------------------------------------------------
# submission monitor
# ---------------------------------------------------------------------------

@bp.route("/packs/<key>/load", methods=["POST"])
@admin_required
def pack_load(key):
    """Fill a department from its Drive documents, already turned into portal data."""
    try:
        report = load_pack(key, _year(), _actor(), replace=request.form.get("replace") == "on")
    except KeyError:
        abort(404)
    dept, s = report["dept"], report["summary"]
    flash(f"{dept['dept_name']}: {len(report['written'])} programme part(s) filled from the Drive "
          f"folder — {s['submitted']} submitted, {s['draft']} saved as drafts listing what is "
          f"still to complete. {len(report['attached'])} file(s) attached"
          + (f", {len(report['missing'])} not in the pack." if report["missing"] else ".")
          + (f" {len(report['skipped'])} part(s) left as they were." if report["skipped"] else ""),
          "success")
    if report["password"]:
        return redirect(url_for("admin.credential_slip", dept_code=dept["dept_code"]))
    return redirect(url_for("admin.submission_detail", dept_code=dept["dept_code"]))


@bp.route("/submissions")
@admin_required
def submissions():
    """Overview → Monitor: every department, stage by stage."""
    data = analysis_sheets(_year())
    rows = sorted(data["stages"], key=lambda r: (-sum(c["status"] == "submitted" for c in r["cells"]),
                                                 r["dept"].get("dept_name", "")))
    return render_template("admin/submissions.html", rows=rows, year=_year(), STAGES=STAGES)


@bp.route("/submissions/<dept_code>")
@admin_required
def submission_detail(dept_code):
    db = get_db()
    year = _year()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    sub = get_or_create_submission(dept_code, year)
    # from Department Information once saved, else the catalogue it starts from
    prog_names = {p["programme_code"]: p.get("programme_name") or p["programme_code"]
                  for p in programmes_of(sub, dept)}
    return render_template("admin/submission_detail.html", dept=dept, submission=sub,
                           prog_names=prog_names,
                           board=stage_board(sub), progress=progress(sub),
                           year=year, STAGE_BY_KEY=STAGE_BY_KEY,
                           documents=_documents(dept_code, sub, year),
                           batches=batches(),
                           summaries=True)


# ---------------------------------------------------------------------------
# a department's uploaded documents: view them, and a short AI summary of each PDF
# ---------------------------------------------------------------------------

def _stored_names(node, out):
    """Every uploaded file still attached somewhere in the submission."""
    if isinstance(node, dict):
        if node.get("stored") and node.get("name"):
            out.append(node["stored"])
        for v in node.values():
            _stored_names(v, out)
    elif isinstance(node, list):
        for v in node:
            _stored_names(v, out)
    return out


def _documents(dept_code, sub, year):
    from .summarise import field_label
    names = list(dict.fromkeys(_stored_names(sub.get("stages", {}), []) +
                               _stored_names(sub.get("programmes", {}), [])))
    if not names:
        return []
    recs = {r["stored_name"]: r for r in get_db().files.find(
        {"dept_code": dept_code, "stored_name": {"$in": names}})}
    groups, order = {}, []
    for n in names:
        r = recs.get(n)
        if not r:
            continue
        stage = STAGE_BY_KEY.get(r.get("stage"), {})
        t, grp = stage.get("title", r.get("stage", "")), stage.get("group", "")
        title = grp if t and t in grp else " · ".join(x for x in (grp, t) if x)
        if title not in groups:
            groups[title] = []
            order.append(title)
        url = url_for("admin.document", dept_code=dept_code, stored=n)
        is_pdf = r["original_name"].lower().endswith(".pdf")
        groups[title].append({
            "name": r["original_name"], "size": r.get("size"),
            "label": field_label(r.get("stage", ""), r.get("field", "")),
            "url": url, "is_pdf": is_pdf,
            "viewable": is_pdf or r["original_name"].lower().rsplit(".", 1)[-1] in ("png", "jpg", "jpeg", "webp", "gif"),
            "thumb": url + "?thumb=1" if is_pdf else None,
            "summary": r.get("summary"), "summary_by": r.get("summary_by") or "AI",
            "uploaded_at": r.get("uploaded_at"),
            "match": r.get("keyword_match"),
        })
    return [{"title": t, "files": groups[t]} for t in order]


def _file_rec(dept_code, stored):
    rec = get_db().files.find_one({"dept_code": dept_code, "stored_name": stored}) or abort(404)
    path = (current_app.config["UPLOAD_ROOT"] / (rec.get("academic_year") or _year())
            / dept_code / rec["stage"] / stored)
    return rec, path


@bp.route("/documents/<dept_code>/<stored>")
@admin_required
def document(dept_code, stored):
    from pathlib import Path
    from .dept import INLINE_EXT, _thumbnail
    rec, path = _file_rec(dept_code, stored)
    if not path.exists():
        abort(404)
    suffix = Path(rec["original_name"]).suffix.lower()
    if request.args.get("thumb") == "1":
        drawn = _thumbnail(path) if suffix == ".pdf" else None
        if not drawn:
            abort(404)
        return send_file(drawn, mimetype="image/png", max_age=86400)
    inline = request.args.get("inline") == "1" and suffix in INLINE_EXT
    return send_file(path, as_attachment=not inline, download_name=rec["original_name"])


@bp.post("/documents/<dept_code>/<stored>/summary")
@admin_required
def document_summary(dept_code, stored):
    from .summarise import SummaryError, summarise
    rec, path = _file_rec(dept_code, stored)
    try:
        out = summarise(rec, path, refresh=request.args.get("refresh") == "1")
    except SummaryError as e:
        return jsonify({"ok": False, "error": str(e)})
    return jsonify({"ok": True, **out})


@bp.route("/submissions/<dept_code>/<stage_key>/return", methods=["POST"])
@admin_required
def submission_return(dept_code, stage_key):
    note = (request.form.get("note") or "").strip()
    if not note:
        flash("Say what needs correcting before you send a stage back.", "error")
        return redirect(url_for("admin.submission_detail", dept_code=dept_code))
    prog = (request.form.get("programme") or "").strip()
    stage = STAGE_BY_KEY.get(stage_key) or abort(404)
    if stage.get("parent") and prog:
        # one programme's part, on its own
        path = f"programmes.{prog}.{stage_key}"
        get_db().submissions.update_one(
            {"dept_code": dept_code, "academic_year": _year()},
            {"$set": {f"{path}.status": "returned", f"{path}.returned_note": note,
                      f"{path}.returned_by": _actor(), f"{path}.returned_at": now(),
                      "status": "in_progress", "updated_at": now()}})
        target = f"{dept_code}/{prog}/{stage_key}"
    else:
        return_stage(dept_code, _year(), stage_key, note, _actor())
        target = f"{dept_code}/{stage_key}"
    audit(_actor(), "stage.returned", target, {"note": note})
    flash(f"“{stage['title']}”{' (' + prog + ')' if prog else ''} has been sent back for correction.", "info")
    return redirect(url_for("admin.submission_detail", dept_code=dept_code))


@bp.route("/submissions/<dept_code>/review/<stage_key>")
@admin_required
def submission_review(dept_code, stage_key):
    """Everything a department entered in one stage (or one programme's
    part), for the review pop-up: the form's own layout, files pointed at
    the Office's copies, how much is filled and how long it took."""
    from .workflow import stage_timing
    stage = STAGE_BY_KEY.get(stage_key) or abort(404)
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    sub = get_or_create_submission(dept_code, _year())
    prog = request.args.get("programme") or ""
    if prog:
        state = (((sub.get("programmes") or {}).get(prog) or {}).get(stage_key) or {})
        status = state.get("status") or "open"
    else:
        state = (sub.get("stages") or {}).get(stage_key) or {}
        status = compute_status(sub, stage_key)

    def point(node):
        # a file's link becomes the admin's link to the same file
        if isinstance(node, dict):
            if node.get("stored") and node.get("name"):
                node = dict(node)
                node["url"] = url_for("admin.document", dept_code=dept_code, stored=node["stored"])
                return node
            return {k: point(v) for k, v in node.items()}
        if isinstance(node, list):
            return [point(v) for v in node]
        return node

    pname = ""
    if prog:
        pname = next((p.get("programme_name", "") for p in
                      ((sub.get("stages") or {}).get("dept_info", {}).get("data", {}) or {}).get("programmes_offered", [])
                      if p.get("programme_code") == prog), "") or prog
    return jsonify({
        "stage": {"key": stage_key, "title": stage.get("title"), "sections": stage.get("sections", [])},
        "data": point(state.get("data") or {}),
        "status": status,
        "timing": stage_timing(state),
        "summary": state.get("summary") or {},
        "returned_note": state.get("returned_note") or "",
        "dept_name": dept.get("dept_name", dept_code), "dept_code": dept_code,
        "comment_url": url_for("admin.comment_add"),
        "comments": [{"text": c["text"], "at": c["at"].strftime("%d %b %Y"), "status": c.get("status"),
                      "reply": c.get("reply", "")}
                     for c in db.comments.find({"dept_code": dept_code, "academic_year": _year(),
                                                "stage": stage_key, "programme_code": prog}).sort("at", -1)],
        "programme": prog, "programme_name": pname,
        "return_url": url_for("admin.submission_return", dept_code=dept_code, stage_key=stage_key),
        "unlock_url": url_for("admin.submission_unlock", dept_code=dept_code, stage_key=stage_key),
    })


@bp.route("/submissions/<dept_code>/<stage_key>/unlock", methods=["POST"])
@admin_required
def submission_unlock(dept_code, stage_key):
    unlock_stage(dept_code, _year(), stage_key, _actor())
    audit(_actor(), "stage.unlocked", f"{dept_code}/{stage_key}")
    flash(f"“{STAGE_BY_KEY[stage_key]['title']}” has been opened for this department.", "info")
    return redirect(url_for("admin.submission_detail", dept_code=dept_code))


# ---------------------------------------------------------------------------
# rule configuration
# ---------------------------------------------------------------------------

@bp.route("/rules", methods=["GET", "POST"])
@admin_required
def rules():
    db = get_db()
    doc = rules_doc()

    if request.method == "POST":
        table = []
        for row in doc.get("table2", U.DEFAULT_TABLE_2):
            key = row["key"]
            new = {"sl": row["sl"], "key": key, "label": row["label"]}
            for track in ("ug3", "ug4"):
                def val(suffix):
                    raw = request.form.get(f"{key}.{track}.{suffix}", "").strip()
                    return int(raw) if raw.isdigit() else None
                new[track] = {
                    "min": val("min") or 0,
                    "max": val("max"),
                    "applicable": request.form.get(f"{key}.{track}.applicable") == "on",
                }
            table.append(new)

        totals = {
            "ug3": int(request.form.get("total.ug3") or U.DEFAULT_TOTALS["ug3"]),
            "ug4": int(request.form.get("total.ug4") or U.DEFAULT_TOTALS["ug4"]),
        }
        other = dict(doc.get("other") or U.DEFAULT_OTHER_RULES)
        for k in ("marks_per_credit", "max_credits_per_course", "meeting_notice_days",
                  "revision_benchmark_threshold"):
            raw = request.form.get(f"other.{k}", "").strip()
            if raw.isdigit():
                other[k] = int(raw)

        db.rules.update_one({"_id": "ugc"},
                            {"$set": {"table2": table, "totals": totals, "other": other,
                                      "updated_at": now(), "updated_by": _actor()}})
        audit(_actor(), "rules.updated")
        flash("The credit rules have been updated. They apply to every validation from now on.",
              "success")
        return redirect(url_for("admin.rules"))

    return render_template("admin/rules.html", doc=doc, tracks=U.TRACKS,
                           defaults=U.DEFAULT_TABLE_2, in_lieu=U.IN_LIEU_RULE)


@bp.route("/rules/reset", methods=["POST"])
@admin_required
def rules_reset():
    get_db().rules.update_one({"_id": "ugc"},
                              {"$set": {"table2": U.DEFAULT_TABLE_2,
                                        "totals": U.DEFAULT_TOTALS,
                                        "other": U.DEFAULT_OTHER_RULES,
                                        "updated_at": now(), "updated_by": _actor()}})
    audit(_actor(), "rules.reset")
    flash("Credit rules restored to the published UGC Table 2 values.", "success")
    return redirect(url_for("admin.rules"))


# ---------------------------------------------------------------------------
# settings, audit, exports
# ---------------------------------------------------------------------------

@bp.route("/settings", methods=["GET", "POST"])
@admin_required
def app_settings():
    db = get_db()
    if request.method == "POST":
        from .db import settings as _s
        was_dev = bool(_s().get("dev_mode"))
        dev = request.form.get("dev_mode") == "on"
        fields = {"academic_year": (request.form.get("academic_year") or "").strip(),
                  "submissions_open": request.form.get("submissions_open") == "on",
                  "dev_mode": dev,
                  "banner": (request.form.get("banner") or "").strip(),
                  "updated_at": now()}
        # the batches: the current one, and the earlier ones whose syllabi
        # departments keep — "2024-2025", one per line (or comma-separated)
        if "current_batch" in request.form:
            from .schema import batch_start
            bad = []
            cur = batch_start(request.form.get("current_batch"))
            if not cur or not 2010 <= cur <= 2040:
                bad.append(request.form.get("current_batch") or "(empty)")
            else:
                fields["current_batch"] = f"{cur}-{cur + 1}"
            years = []
            for raw in re.split(r"[,\n]+", request.form.get("existing_batches") or ""):
                if not raw.strip():
                    continue
                y = batch_start(raw)
                if not y or not 2010 <= y <= 2040:
                    bad.append(raw.strip())
                elif y != cur and f"{y}-{y + 1}" not in years:
                    years.append(f"{y}-{y + 1}")
            fields["existing_batches"] = sorted(years)
            if bad:
                flash("Not a batch year (use the form 2024-2025, from 2010 to 2040): "
                      + ", ".join(bad), "error")
                return redirect(url_for("admin.app_settings"))
        db.settings.update_one({"_id": "app"}, {"$set": fields})
        audit(_actor(), "settings.updated")
        # Unlocking every stage for every department is worth its own line in
        # the log, separate from whatever else was saved in the same form.
        if dev != was_dev:
            audit(_actor(), "settings.dev_mode." + ("on" if dev else "off"))
            flash("Developer mode is ON — every stage is unlocked for every "
                  "department." if dev else
                  "Developer mode is off. Stages lock in sequence again.",
                  "warning" if dev else "success")
        else:
            flash("Settings saved.", "success")
        return redirect(url_for("admin.app_settings"))
    from .db import settings as s
    return render_template("admin/settings.html", settings=s(), batches=batches(),
                           stage_count=len(STAGES))


@bp.route("/audit")
@admin_required
def audit_log():
    entries = list(get_db().audit.find().sort("at", -1).limit(300))
    return render_template("admin/audit.html", entries=entries)


@bp.route("/export/institution.xlsx")
@admin_required
def export_institution():
    buf = institution_excel(_year())
    audit(_actor(), "export.institution")
    return send_file(buf, as_attachment=True,
                     download_name=f"BoS-Repository-{_year()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/export/<dept_code>.xlsx")
@admin_required
def export_department(dept_code):
    buf = department_excel(dept_code, _year(), ShareLinks(dept_code, _year()))
    audit(_actor(), "export.department", dept_code)
    return send_file(buf, as_attachment=True,
                     download_name=f"BoS-{dept_code}-{_year()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/report/<dept_code>/<programme_code>")
@admin_required
def programme_report(dept_code, programme_code):
    """One programme's whole record as a page to read or print."""
    from .report import programme_context
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    ctx = programme_context(db, dept, _year(), programme_code, AdminLinks(dept_code, _year())) or abort(404)
    return render_template("admin/programme_report.html", **ctx)


@bp.route("/export/<dept_code>/preview/<kind>")
@admin_required
def export_preview(dept_code, kind):
    """The Excel or Word download, shown in a browser tab: every sheet as a
    tab of its own, with the download beside it."""
    from .share import _docx_blocks, _xlsx_sheets
    dept = get_db().departments.find_one({"dept_code": dept_code}) or abort(404)
    links = ShareLinks(dept_code, _year())
    if kind == "xlsx":
        buf, body = department_excel(dept_code, _year(), links), None
        body = _xlsx_sheets(buf, max_rows=3000, max_cols=12)
        name, url = f"BoS-{dept_code}-{_year()}.xlsx", url_for("admin.export_department", dept_code=dept_code)
    elif kind == "docx":
        buf = submission_word(dept_code, _year(), links)
        body = _docx_blocks(buf)
        name, url = f"BoS-Report-{dept_code}-{_year()}.docx", url_for("admin.export_department_word",
                                                                      dept_code=dept_code)
    else:
        abort(404)
    rec = {"original_name": name, "dept_code": dept_code, "academic_year": _year(),
           "size": len(buf.getvalue())}
    return render_template("share/doc.html", rec=rec, dept=dept, kind=kind, body=body,
                           file_url=url, download_url=url, hide_chrome=True)


@bp.post("/departments/<dept_code>/wipe")
@admin_required
def wipe_department(dept_code):
    """Delete everything the department has entered and uploaded — every year,
    every stage, every file — keeping the department and its login."""
    import shutil
    db = get_db()
    dept = db.departments.find_one({"dept_code": dept_code}) or abort(404)
    if (request.form.get("confirm") or "").strip().upper() != dept_code.upper():
        flash(f"Type the department code {dept_code} to confirm. Nothing was deleted.", "error")
        return redirect(url_for("admin.submission_detail", dept_code=dept_code))
    counts = {}
    for col in ("submissions", "files", "comments", "notifications", "versions"):
        counts[col] = db[col].delete_many({"dept_code": dept_code}).deleted_count
    root = current_app.config["UPLOAD_ROOT"]
    for year_dir in (root.iterdir() if root.is_dir() else []):
        folder = year_dir / dept_code
        if folder.is_dir():
            shutil.rmtree(folder, ignore_errors=True)
    audit(_actor(), "department.wiped", dept_code, counts)
    flash(f"All of {dept['dept_name']}'s data is deleted: {counts['submissions']} record(s), "
          f"{counts['files']} file(s). The department and its login are kept.", "success")
    return redirect(url_for("admin.submission_detail", dept_code=dept_code))


@bp.route("/report/<dept_code>/<programme_code>/<kind>.docx")
@admin_required
def programme_docx(dept_code, programme_code, kind):
    """The programme's Curriculum ("c") or Syllabus ("s") in the Office's templates."""
    from .share import DOCX, generated_docx
    if kind not in ("c", "s", "v"):
        abort(404)
    out = generated_docx(kind, dept_code, programme_code, _year()) or abort(404)
    return send_file(out[0], as_attachment=True, download_name=out[1], mimetype=DOCX)


@bp.route("/export/<dept_code>.docx")
@admin_required
def export_department_word(dept_code):
    buf = submission_word(dept_code, _year(), ShareLinks(dept_code, _year()))
    audit(_actor(), "export.word", dept_code)
    return send_file(buf, as_attachment=True,
                     download_name=f"BoS-Report-{dept_code}-{_year()}.docx",
                     mimetype="application/vnd.openxmlformats-officedocument."
                              "wordprocessingml.document")


# ---------------------------------------------------------------------------
# updates: what departments changed, as it happens
# ---------------------------------------------------------------------------

def _updates_query(args):
    q = {}
    if args.get("dept"):
        q["dept_code"] = args["dept"]
    if args.get("event"):
        q["event"] = args["event"]
    if args.get("show") == "unread":
        q["read"] = False
    return q


@bp.route("/updates")
@admin_required
def updates():
    from .notify import EVENTS
    db = get_db()
    q = _updates_query(request.args)
    items = list(db.notifications.find(q).sort("at", -1).limit(300))
    days, order = {}, []
    for n in items:
        d = n["at"].strftime("%A, %d %B %Y")
        if d not in days:
            days[d] = []
            order.append(d)
        days[d].append(n)
    depts = sorted({(n["dept_code"], n.get("dept_name", "")) for n in
                    db.notifications.find({}, {"dept_code": 1, "dept_name": 1})}, key=lambda x: x[1])
    counts = {k: db.notifications.count_documents({"event": k}) for k in EVENTS}
    return render_template("admin/updates.html", days=[(d, days[d]) for d in order],
                           events=EVENTS, depts=depts, counts=counts, args=request.args,
                           unread=db.notifications.count_documents({"read": False}),
                           total=db.notifications.count_documents({}))


@bp.post("/updates/read")
@admin_required
def updates_read():
    from bson import ObjectId
    db = get_db()
    one = request.form.get("id")
    if one:
        try:
            db.notifications.update_one({"_id": ObjectId(one)}, {"$set": {"read": True}})
        except Exception:
            abort(400)
    else:
        db.notifications.update_many({"read": False}, {"$set": {"read": True}})
    if request.accept_mimetypes.best == "application/json" or request.is_json:
        return jsonify({"ok": True})
    return redirect(request.referrer or url_for("admin.updates"))


@bp.route("/updates.json")
@admin_required
def updates_json():
    """For the bell: how many are unread, and any newer than `after`."""
    db = get_db()
    q = {"read": False}
    after = request.args.get("after")
    items = []
    if after:
        try:
            since = datetime.fromisoformat(after)
            items = [{"id": str(n["_id"]), "text": n["text"], "event": n["event"],
                      "at": n["at"].isoformat(),
                      "link": url_for("admin.submission_detail", dept_code=n["dept_code"])}
                     for n in db.notifications.find({"at": {"$gt": since}}).sort("at", 1).limit(5)]
        except ValueError:
            pass
    return jsonify({"unread": db.notifications.count_documents(q), "items": items,
                    "now": now().isoformat()})


# ---------------------------------------------------------------------------
# connectors: where updates are sent
# ---------------------------------------------------------------------------

def _connector_form(f):
    from .notify import EVENTS, KINDS
    kind = f.get("kind")
    if kind not in KINDS:
        return None, "Choose what kind of connector this is."
    target = (f.get("target") or "").strip()
    if kind == "email":
        emails = [e.strip() for e in target.replace(";", ",").split(",") if e.strip()]
        if not emails or not all(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", e) for e in emails):
            return None, "Give one or more email addresses, separated by commas."
        target = ", ".join(emails)
    elif not target.startswith("https://"):
        return None, "The webhook address must start with https://"
    events = [e for e in f.getlist("events") if e in EVENTS]
    if not events:
        return None, "Choose at least one kind of update to send."
    return {"kind": kind, "name": (f.get("name") or KINDS[kind]).strip()[:80], "target": target,
            "events": events, "departments": (f.get("departments") or "").strip()[:400],
            "active": True}, None


def _mask(c):
    if c["kind"] == "email":
        return c["target"]
    t = c["target"]
    return t[:28] + "…" + t[-4:] if len(t) > 40 else t


@bp.route("/connectors", methods=["GET", "POST"])
@admin_required
def connectors():
    from .notify import EVENTS, KINDS, smtp_ready
    db = get_db()
    if request.method == "POST":
        doc, err = _connector_form(request.form)
        if err:
            flash(err, "error")
        else:
            doc.update(created_at=now(), sent=0, failed=0)
            db.connectors.insert_one(doc)
            audit(_actor(), "connector.added", doc["kind"], {"name": doc["name"]})
            flash(f"{doc['name']} connected. Use “Send test” to check it.", "success")
        return redirect(url_for("admin.connectors"))
    items = list(db.connectors.find().sort("created_at", 1))
    for c in items:
        c["shown"] = _mask(c)
    return render_template("admin/connectors.html", items=items, kinds=KINDS, events=EVENTS,
                           smtp=smtp_ready())


def _connector(cid):
    from bson import ObjectId
    try:
        c = get_db().connectors.find_one({"_id": ObjectId(cid)})
    except Exception:
        c = None
    if not c:
        abort(404)
    return c


@bp.post("/connectors/<cid>/toggle")
@admin_required
def connector_toggle(cid):
    c = _connector(cid)
    get_db().connectors.update_one({"_id": c["_id"]}, {"$set": {"active": not c.get("active")}})
    audit(_actor(), "connector." + ("paused" if c.get("active") else "resumed"), c["kind"])
    return redirect(url_for("admin.connectors"))


@bp.post("/connectors/<cid>/delete")
@admin_required
def connector_delete(cid):
    c = _connector(cid)
    get_db().connectors.delete_one({"_id": c["_id"]})
    audit(_actor(), "connector.removed", c["kind"], {"name": c.get("name")})
    flash(f"{c.get('name')} removed.", "success")
    return redirect(url_for("admin.connectors"))


@bp.post("/connectors/<cid>/test")
@admin_required
def connector_test(cid):
    from .notify import send_one, test_note
    c = _connector(cid)
    ok, words = send_one(c, test_note())
    get_db().connectors.update_one({"_id": c["_id"]}, {
        "$set": {"last_at": now(), "last_ok": ok, "last_status": "Test: " + words}})
    flash(f"Test sent to {c.get('name')}." if ok else f"{c.get('name')}: the test failed — {words}",
          "success" if ok else "error")
    return redirect(url_for("admin.connectors"))


# ---------------------------------------------------------------------------
# keywords: what each upload box is checked for
# ---------------------------------------------------------------------------

@bp.route("/keywords", methods=["GET", "POST"])
@admin_required
def keywords():
    from .keyword_match import default_keywords, overrides, upload_boxes
    db = get_db()
    boxes = upload_boxes()
    if request.method == "POST":
        mine = {}
        for b in boxes:
            raw = request.form.get("kw_" + b["field"], "")
            words = list(dict.fromkeys(w.strip() for w in re.split(r"[,\n]+", raw) if w.strip()))[:30]
            if words and words != default_keywords(b["field"], b["label"]):
                mine[b["field"]] = words
        db.settings.update_one({"_id": "app"}, {"$set": {"keywords": mine}}, upsert=True)
        audit(_actor(), "keywords.updated", detail={"boxes": len(mine)})
        flash("Keywords saved. They apply to the next upload — use “Check all files again” "
              "for the ones already in.", "success")
        return redirect(url_for("admin.keywords"))

    mine = overrides()
    for b in boxes:
        b["default"] = default_keywords(b["field"], b["label"])
        b["words"] = mine.get(b["field"]) or b["default"]
        b["custom"] = b["field"] in mine
        stats = {"match": 0, "weak": 0, "miss": 0, "unread": 0}
        for r in db.files.find({"field": b["field"], "academic_year": _year()}, {"keyword_match": 1}):
            s = (r.get("keyword_match") or {}).get("status")
            if s in stats:
                stats[s] += 1
        b["stats"] = stats

    status = request.args.get("status", "")
    q = {"academic_year": _year(), "keyword_match": {"$exists": True}}
    if status:
        q["keyword_match.status"] = status
    names = {d["dept_code"]: d.get("dept_name", "") for d in db.departments.find({}, {"dept_code": 1, "dept_name": 1})}
    files = []
    for r in db.files.find(q).sort("uploaded_at", -1).limit(400):
        files.append({"dept_code": r["dept_code"], "dept_name": names.get(r["dept_code"], r["dept_code"]),
                      "name": r["original_name"], "field": r.get("field"),
                      "match": r.get("keyword_match") or {}, "at": r.get("uploaded_at"),
                      "url": url_for("admin.document", dept_code=r["dept_code"], stored=r["stored_name"])})
    totals = {k: db.files.count_documents({"academic_year": _year(), "keyword_match.status": k})
              for k in ("match", "weak", "miss", "unread")}
    unchecked = db.files.count_documents({"academic_year": _year(), "keyword_match": {"$exists": False}})
    return render_template("admin/keywords.html", boxes=boxes, files=files, totals=totals,
                           status=status, unchecked=unchecked)


@bp.post("/keywords/reset")
@admin_required
def keywords_reset():
    get_db().settings.update_one({"_id": "app"}, {"$set": {"keywords": {}}}, upsert=True)
    audit(_actor(), "keywords.reset")
    flash("Every box is back to its standard keywords.", "success")
    return redirect(url_for("admin.keywords"))


@bp.post("/keywords/recheck")
@admin_required
def keywords_recheck():
    """Run the check again on every file this year — for files uploaded
    before the check existed, or after the keywords changed."""
    from .keyword_match import check
    from .summarise import field_label
    db = get_db()
    root = current_app.config["UPLOAD_ROOT"]
    done = gone = 0
    for r in db.files.find({"academic_year": _year()}):
        path = root / (r.get("academic_year") or _year()) / r["dept_code"] / r["stage"] / r["stored_name"]
        if not path.exists():
            gone += 1
            continue
        m = check(path, r.get("field", ""), field_label(r.get("stage", ""), r.get("field", "")))
        db.files.update_one({"_id": r["_id"]}, {"$set": {"keyword_match": m}})
        done += 1
    audit(_actor(), "keywords.rechecked", detail={"files": done})
    flash(f"Checked {done} file{'s' if done != 1 else ''} again"
          + (f"; {gone} no longer on the server." if gone else "."), "success")
    return redirect(url_for("admin.keywords"))


# ---------------------------------------------------------------------------
# the demo department
# ---------------------------------------------------------------------------

@bp.post("/demo-department")
@admin_required
def demo_department():
    from . import demo_dept
    try:
        user, pw = demo_dept.create(_year(), current_app.config["UPLOAD_ROOT"], actor=_actor())
    except Exception as e:                       # say why, rather than a bare 500
        current_app.logger.exception("Could not make the demo department")
        flash(f"The demo department could not be made — {e.__class__.__name__}: {e}"[:400], "error")
        return redirect(url_for("admin.departments"))
    audit(_actor(), "demo.created", demo_dept.CODE)
    flash(f"{demo_dept.NAME} is ready. Username {user}, password {pw} — sign in as it "
          "(in a private window) and press “Fill everything with sample data”.", "success")
    return redirect(url_for("admin.departments"))


@bp.post("/demo-department/remove")
@admin_required
def demo_department_remove():
    from . import demo_dept
    demo_dept.remove()
    audit(_actor(), "demo.removed", demo_dept.CODE)
    flash(f"{demo_dept.NAME} removed.", "success")
    return redirect(url_for("admin.departments"))


# ---------------------------------------------------------------------------
# analysis sheets: every department, stage by stage, programme by programme,
# document by document — and comments asking for a change
# ---------------------------------------------------------------------------

BOS_DOCS = [("pre_bos", "pre_bos_files", "diac_signed", "DIAC"), ("pre_bos", "pre_bos_files", "dpac_signed", "DPAC"),
            ("bos_documents", "bos_files", "bos_composition", "BoS composition"),
            ("bos_documents", "bos_files", "external_profiles", "Expert profiles"),
            ("bos_documents", "bos_files", "minutes", "MoM"),
            ("bos_documents", "bos_files", "geotagged_photos", "Photos"),
            ("bos_documents", "bos_files", "attendance", "Attendance"),
            ("bos_documents", "bos_files", "vision_mission", "Vision-Mission"),
            ("bos_documents", "bos_files", "feedback_curriculum", "Feedback"),
            ("bos_documents", "bos_files", "feedback_new_programme", "Feedback (new)")]


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _doc_state(val):
    """One upload box, in a word: '' (nothing), ok, warn or bad, and why."""
    vals = [v for v in (val if isinstance(val, list) else [val]) if isinstance(v, dict) and v.get("name")]
    if not vals:
        return "", ""
    why = []
    worst = "ok"
    for v in vals:
        m = v.get("match") or {}
        t = m.get("template") or {}
        if t and (t.get("blank") or t.get("half") or t.get("placeholders")):
            worst = "bad"
            why.append("blank: " + ", ".join(t.get("blank") or t.get("half") or ["Words only"]))
        elif m.get("looks_like"):
            worst = "bad"
            why.append("looks like " + m["looks_like"])
        elif m.get("status") == "miss":
            worst = "bad" if worst != "bad" else worst
            why.append("keywords do not match")
        elif m.get("status") == "weak" and worst == "ok":
            worst = "warn"
            why.append("few keywords")
    return worst, "; ".join(why) or f"{len(vals)} file{'s' if len(vals) != 1 else ''}"


def analysis_sheets(year):
    from .workflow import programme_stage_state, programmes_of, stage_timing
    db = get_db()
    depts = list(db.departments.find({"active": True}).sort("dept_name", 1))
    subs = {s["dept_code"]: s for s in db.submissions.find({"academic_year": year})}
    open_comments = {}
    for c in db.comments.find({"academic_year": year, "status": "open"}):
        open_comments[c["dept_code"]] = open_comments.get(c["dept_code"], 0) + 1
    stages, progs, docs = [], [], []
    for d in depts:
        sub = subs.get(d["dept_code"]) or {"stages": {}, "programmes": {}}
        row = {"dept": d, "cells": [], "comments": open_comments.get(d["dept_code"], 0)}
        for s in STAGES:
            st = (sub.get("stages") or {}).get(s["key"]) or {}
            row["cells"].append({"key": s["key"], "title": s["title"],
                                 "status": compute_status(sub, s["key"]) if sub.get("_id") else "open",
                                 "timing": stage_timing(st)})
        stages.append(row)
        for p in programmes_of(sub, d) if sub.get("_id") else []:
            code = p["programme_code"]
            part = lambda k: programme_stage_state(sub, code, k)
            cur = (part("prog_curriculum").get("data") or {})
            rows = [r for r in (cur.get("semester_structure") or []) if isinstance(r, dict)]
            syl = (part("prog_syllabus").get("data") or {}).get("courses") or []
            rev = [r for r in ((part("prog_revision").get("data") or {}).get("courses") or []) if isinstance(r, dict)]
            changes = [_num(r.get("avg_change")) for r in rev if r.get("avg_change") not in (None, "")]
            progs.append({"dept": d, "code": code, "name": p.get("programme_name") or code,
                          "degree": (cur.get("details") or {}).get("degree_level") or p.get("degree", ""),
                          "courses": len(rows), "credits": int(sum(_num(r.get("credits")) for r in rows)),
                          "syllabi": len(syl), "revised": len(rev),
                          "avg_change": round(sum(changes) / len(changes), 1) if changes else None,
                          "parts": {k: part(k).get("status") or "open"
                                    for k in ("prog_curriculum", "prog_syllabus", "prog_revision")}})
        drow = {"dept": d, "cells": []}
        for stage, sec, field, label in BOS_DOCS:
            val = (((sub.get("stages") or {}).get(stage) or {}).get("data") or {}).get(sec, {})
            state, why = _doc_state((val or {}).get(field))
            drow["cells"].append({"label": label, "state": state, "why": why, "stage": stage})
        docs.append(drow)
    comments = list(db.comments.find({"academic_year": year}).sort("at", -1).limit(500))
    names = {d["dept_code"]: d.get("dept_name", "") for d in depts}
    return {"stages": stages, "programmes": progs, "documents": docs, "comments": comments,
            "names": names, "doc_labels": [x[3] for x in BOS_DOCS]}


@bp.route("/sheets")
@admin_required
def sheets():
    tab = request.args.get("tab", "programmes")
    if tab not in ("programmes", "documents", "comments"):
        return redirect(url_for("admin.submissions"))
    data = analysis_sheets(_year())
    return render_template("admin/sheets.html", year=_year(), tab=tab,
                           STAGES=STAGES, STAGE_BY_KEY=STAGE_BY_KEY, **data)


@bp.route("/sheets.xlsx")
@admin_required
def sheets_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = analysis_sheets(_year())
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0E2A5C")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))
        ws.freeze_panes = "B2"

    ws = wb.active
    ws.title = "Stages"
    sheet(ws, ["Department", "Code", "Campus"] + [f"{s['title']}" for s in STAGES] + ["Open comments"],
          [[r["dept"].get("dept_name"), r["dept"]["dept_code"], r["dept"].get("campus", "")]
           + [c["status"] + (f" ({c['timing']['took']})" if c["timing"].get("took") else "") for c in r["cells"]]
           + [r["comments"]] for r in data["stages"]])
    sheet(wb.create_sheet("Programmes"),
          ["Department", "Programme", "Code", "Degree", "Courses", "Credits", "Syllabi", "Revised courses",
           "Avg % change", "Curriculum", "Current syllabus", "Revision"],
          [[p["dept"].get("dept_name"), p["name"], p["code"], p["degree"], p["courses"], p["credits"],
            p["syllabi"], p["revised"], p["avg_change"], p["parts"]["prog_curriculum"],
            p["parts"]["prog_syllabus"], p["parts"]["prog_revision"]] for p in data["programmes"]])
    sheet(wb.create_sheet("Documents"), ["Department"] + data["doc_labels"],
          [[r["dept"].get("dept_name")] + [(c["state"] or "—") + (f": {c['why']}" if c["state"] in ("warn", "bad") else "")
                                         for c in r["cells"]] for r in data["documents"]])
    sheet(wb.create_sheet("Comments"), ["When", "Department", "Stage", "Programme", "Comment", "By", "Status", "Reply"],
          [[c["at"].strftime("%d %b %Y %H:%M"), data["names"].get(c["dept_code"], c["dept_code"]),
            STAGE_BY_KEY.get(c["stage"], {}).get("title", c["stage"]), c.get("programme_code", ""),
            c["text"], c.get("by", ""), c.get("status", ""), c.get("reply", "")] for c in data["comments"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    audit(_actor(), "export.sheets")
    return send_file(buf, as_attachment=True, download_name=f"BoS-Analysis-{_year()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/comments")
@admin_required
def comment_add():
    """A comment asking a department to change something — on a stage or
    a programme's part. If it is submitted, it can be sent back with it."""
    db = get_db()
    f = request.form
    code, stage = (f.get("dept") or "").strip(), (f.get("stage") or "").strip()
    text = (f.get("text") or "").strip()
    dept = db.departments.find_one({"dept_code": code})
    if not dept or stage not in STAGE_BY_KEY or not text:
        flash("Choose a department and a stage, and write the comment.", "error")
        return redirect(request.referrer or url_for("admin.sheets"))
    prog = (f.get("programme") or "").strip()
    doc = {"dept_code": code, "academic_year": _year(), "stage": stage, "programme_code": prog,
           "text": text[:2000], "by": _actor(), "at": now(), "status": "open", "reopened": False}
    if f.get("reopen") == "on":
        from .workflow import programme_stage_state
        sub = get_or_create_submission(code, _year())
        if prog:
            if programme_stage_state(sub, prog, stage).get("status") == "submitted":
                path = f"programmes.{prog}.{stage}"
                db.submissions.update_one({"_id": sub["_id"]}, {"$set": {
                    f"{path}.status": "returned", f"{path}.returned_note": text,
                    f"{path}.returned_by": _actor(), f"{path}.returned_at": now(), "status": "in_progress"}})
                doc["reopened"] = True
        elif compute_status(sub, stage) == "submitted":
            return_stage(code, _year(), stage, text, _actor())
            doc["reopened"] = True
    db.comments.insert_one(doc)
    audit(_actor(), "comment.added", f"{code}/{prog + '/' if prog else ''}{stage}", {"reopened": doc["reopened"]})
    flash(f"Comment sent to {dept.get('dept_name')}" + (" — the stage is open again for them to change it."
                                                         if doc["reopened"] else "."), "success")
    return redirect(request.referrer or url_for("admin.sheets", tab="comments"))


@bp.post("/comments/<cid>/close")
@admin_required
def comment_close(cid):
    from bson import ObjectId
    try:
        get_db().comments.update_one({"_id": ObjectId(cid)}, {"$set": {"status": "closed", "closed_at": now()}})
    except Exception:
        abort(400)
    return redirect(request.referrer or url_for("admin.sheets", tab="comments"))


@bp.route("/comments.json")
@admin_required
def comments_json():
    q = {"academic_year": _year(), "dept_code": request.args.get("dept", "")}
    if request.args.get("stage"):
        q["stage"] = request.args["stage"]
    if request.args.get("programme") is not None:
        q["programme_code"] = request.args.get("programme", "")
    return jsonify([{"text": c["text"], "by": c.get("by"), "at": c["at"].strftime("%d %b %Y %H:%M"),
                     "status": c.get("status"), "reply": c.get("reply", "")}
                    for c in get_db().comments.find(q).sort("at", -1)])


# ---------------------------------------------------------------------------
# calendar: what happened on which day
# ---------------------------------------------------------------------------

@bp.route("/calendar")
@admin_required
def calendar():
    """A month of activity: a dot on each day something happened, coloured by
    what; click a day for everything done on it."""
    import calendar as cal
    from datetime import date, timedelta
    from .notify import EVENTS
    db = get_db()
    today = now().date()
    try:
        y, m = (int(x) for x in (request.args.get("month") or today.strftime("%Y-%m")).split("-"))
        first = date(y, m, 1)
    except (ValueError, TypeError):
        first = today.replace(day=1)
    last = first.replace(day=cal.monthrange(first.year, first.month)[1])
    start = datetime(first.year, first.month, 1)
    end = datetime(last.year, last.month, last.day) + timedelta(days=1)

    days = {}
    for n in db.notifications.find({"at": {"$gte": start, "$lt": end}}, {"at": 1, "event": 1}):
        d = days.setdefault(n["at"].date(), {"total": 0, "events": {}})
        d["total"] += 1
        d["events"][n["event"]] = d["events"].get(n["event"], 0) + 1
    for c in db.comments.find({"at": {"$gte": start, "$lt": end}}, {"at": 1}):
        d = days.setdefault(c["at"].date(), {"total": 0, "events": {}})
        d["total"] += 1
        d["events"]["comment"] = d["events"].get("comment", 0) + 1

    weeks = cal.Calendar(firstweekday=0).monthdatescalendar(first.year, first.month)
    try:
        chosen = date.fromisoformat(request.args.get("day") or "")
    except ValueError:
        chosen = today if first <= today <= last else None
    items, comments = [], []
    if chosen:
        a = datetime(chosen.year, chosen.month, chosen.day)
        items = list(db.notifications.find({"at": {"$gte": a, "$lt": a + timedelta(days=1)}}).sort("at", -1))
        comments = list(db.comments.find({"at": {"$gte": a, "$lt": a + timedelta(days=1)}}).sort("at", -1))
    prev_m = (first - timedelta(days=1)).strftime("%Y-%m")
    next_m = (last + timedelta(days=1)).strftime("%Y-%m")
    names = {d["dept_code"]: d.get("dept_name", "") for d in db.departments.find({}, {"dept_code": 1, "dept_name": 1})}
    return render_template("admin/calendar.html", weeks=weeks, first=first, days=days, today=today,
                           chosen=chosen, items=items, comments=comments, prev_m=prev_m, next_m=next_m,
                           events=EVENTS, names=names, STAGE_BY_KEY=STAGE_BY_KEY)
