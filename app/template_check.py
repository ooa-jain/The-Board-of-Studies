"""Is a signed composition form filled in?

The university's two templates — Composition of DIAC, and Composition of
the Program Assessment Committee (PAC, the department's DPAC) — are a table:
Category, Role, Name, Designation, with “Words only” written where a name
goes. When a department uploads its filled form, each category is checked
for at least one name and designation; a category left empty, or still
saying “Words only”, is reported as blank, by name, so the department knows
exactly what to fill.

A Word file is read cell by cell. A PDF has no cells to read, so for a PDF
the check is that every category is mentioned and no “Words only” is left.
A scanned PDF has no text at all and is not judged.
"""

from __future__ import annotations

import re
from pathlib import Path

PLACEHOLDER = re.compile(r"^\s*(words\s*only|name|designation|-+|—|nil|\.+)?\s*$", re.I)
# “NA” is a deliberate answer — the department has no such member — not a blank
NOT_APPLICABLE = re.compile(r"^\s*(na|n/a|n\.a\.?|not\s+applicable)\s*$", re.I)

# each template's categories, in its order: (how the form names it, how to find it)
TEMPLATES = {
    "diac_signed": {
        "title": "Composition of DIAC",
        "file": "templates/Composition_of_DIAC.docx",
        "categories": [
            ("Dean of Faculty / Director of School", r"dean|director"),
            ("Head of Department", r"head\s+of\s+(the\s+)?department|\bhod\b"),
            ("Area Chair / Area Head", r"area\s+(chair|head)"),
            ("Program Head / Program Co-Ordinator", r"program(me)?\s+(head|co-?ordinator)"),
            ("Faculty Placement Coordinators", r"placement"),
            ("Industry Experts", r"industry"),
            ("Alumni", r"alumn"),
        ],
    },
    "bos_composition": {
        "title": "Composition of BoS Members",
        "file": "templates/Composition_of_BoS_Members.docx",
        "categories": None,          # the PAC's, filled in below
    },
    "dpac_signed": {
        "title": "Composition of the Program Assessment Committee",
        "file": "templates/Composition_of_PAC.docx",
        "categories": [
            ("Dean of Faculty / Director of School", r"dean|director"),
            ("Head of Department", r"head\s+of\s+(the\s+)?department|\bhod\b"),
            ("Area Chair / Area Head", r"area\s+(chair|head)"),
            ("Program Head / Program Co-Ordinator", r"program(me)?\s+(head|co-?ordinator)"),
            ("Senior Non-Teaching Staff", r"non-?\s*teaching"),
            ("Representative of the Office of Academics", r"office\s+of\s+academics|\booa\b"),
            ("Current Student", r"student"),
            ("Professor", r"(?<!associate )(?<!assistant )\bprofessor\b"),
            ("Associate Professor", r"associate\s+professor"),
            ("Assistant Professor", r"assistant\s+professor"),
            ("Industry Expert", r"industry"),
            ("Alumni", r"alumn"),
            ("Parent", r"\bparent"),
            ("Academician", r"academician"),
        ],
    },
}


TEMPLATES["bos_composition"]["categories"] = TEMPLATES["dpac_signed"]["categories"]


def _which(categories, text):
    t = text.lower()
    for name, pat in categories:
        if re.search(pat, t, re.I):
            return name
    return None


def _cell(c):
    return " ".join(c.text.split())


def check_docx(path: Path, tpl: dict) -> dict | None:
    import docx
    doc = docx.Document(str(path))
    for table in doc.tables:
        rows = table.rows
        head = None
        for i, r in enumerate(rows[:4]):
            texts = [_cell(c).lower() for c in r.cells]
            who = next((t for t in ("name", "members", "member", "name of the member") if t in texts), None)
            if who:
                head = (i, texts.index("category") if "category" in texts else 1,
                        texts.index(who), texts.index("designation") if "designation" in texts else None)
                break
        if not head:
            continue
        start, ci, ni, di = head
        found = {name: {"rows": 0, "filled": 0, "no_designation": 0, "na": 0}
                 for name, _ in tpl["categories"]}
        placeholders = 0
        for r in rows[start + 1:]:
            cells = r.cells
            if len(cells) <= max(ci, ni, di if di is not None else 0):
                continue
            cat = _which(tpl["categories"], " ".join(cells[ci].text.split())[:90])
            if not cat:
                continue
            name = _cell(cells[ni])
            des = _cell(cells[di]) if di is not None else "given"
            f = found[cat]
            f["rows"] += 1
            if re.match(r"^\s*words\s*only\s*$", name, re.I) or re.match(r"^\s*words\s*only\s*$", des, re.I):
                placeholders += 1
            if NOT_APPLICABLE.match(name):
                f["na"] += 1
            elif not PLACEHOLDER.match(name):
                f["filled"] += 1
                if PLACEHOLDER.match(des):
                    f["no_designation"] += 1
        return _verdict(tpl, found, placeholders, "docx")
    return None


def check_text(text: str, tpl: dict) -> dict:
    """For a PDF: every category mentioned, no “Words only” left."""
    found = {}
    for name, pat in tpl["categories"]:
        hit = bool(re.search(pat, text, re.I))
        found[name] = {"rows": int(hit), "filled": int(hit), "no_designation": 0}
    placeholders = len(re.findall(r"words\s*only", text, re.I))
    return _verdict(tpl, found, placeholders, "pdf")


def _verdict(tpl, found, placeholders, kind):
    blank = [n for n, f in found.items() if not f["filled"] and not f.get("na")]
    half = [n for n, f in found.items() if f["filled"] and f["no_designation"]]
    na = [n for n, f in found.items() if not f["filled"] and f.get("na")]
    return {
        "template": tpl["title"], "kind": kind,
        "categories": [{"name": n, **f} for n, f in found.items()],
        "filled": sum(1 for f in found.values() if f["filled"] or f.get("na")),
        "not_applicable": na,
        "total": len(found),
        "blank": blank, "half": half, "placeholders": placeholders,
        "ok": not blank and not half and not placeholders,
    }


def check(path: Path, field: str, text: str | None = None) -> dict | None:
    """The template check for an upload box that has a template, or None."""
    tpl = TEMPLATES.get(field)
    if not tpl:
        return None
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return check_docx(path, tpl)
        if text and len(text.strip()) >= 20:
            return check_text(text, tpl)
    except Exception:
        return None
    return None


def problems(result: dict | None) -> list[str]:
    """What to fill, in words; [] when the form is complete."""
    if not result or result.get("ok"):
        return []
    out = []
    if result.get("blank"):
        out.append("Blank — fill " + ", ".join(result["blank"]))
    if result.get("half"):
        out.append("Designation missing for " + ", ".join(result["half"]))
    if result.get("placeholders"):
        n = result["placeholders"]
        out.append(f"“Words only” is still written in {n} place{'s' if n != 1 else ''} — replace it with the names")
    return out


def fill(src: Path, dst: Path, field: str, skip=(), people=None) -> Path:
    """A copy of a template with a name and designation in the first row of
    each category — for the demo department, and for the tests. Categories
    in `skip` are left as they are."""
    import docx
    tpl = TEMPLATES[field]
    doc = docx.Document(str(src))
    table = doc.tables[0]
    rows = table.rows
    start = next(i for i, r in enumerate(rows[:4]) if "Name" in [_cell(c) for c in r.cells])
    texts = [_cell(c) for c in rows[start].cells]
    ci, ni, di = texts.index("Category"), texts.index("Name"), texts.index("Designation")
    done = set()
    people = people or {}
    for n, r in enumerate(rows[start + 1:], 1):
        cat = _which(tpl["categories"], " ".join(r.cells[ci].text.split())[:90])
        if not cat or cat in skip:
            continue
        name, des = people.get(cat, (f"Dr. Sample Person {n}", cat.split(" /")[0]))
        if cat in done:
            # later rows of a category: clear any “Words only” left in them
            for i in (ni, di):
                if re.match(r"^\s*words\s*only\s*$", _cell(r.cells[i]), re.I):
                    r.cells[i].text = ""
            continue
        r.cells[ni].text = name
        r.cells[di].text = des
        done.add(cat)
    doc.save(str(dst))
    return dst


def make_bos_composition(dst: Path, people: dict) -> Path:
    """A Composition of BoS Members in the department's own layout —
    S.No, Category, Role, Members — one row per category. For the demo."""
    import docx
    roles = {"Dean of Faculty / Director of School": "Chairperson", "Head of Department": "Co-Chairperson"}
    externals = {"Industry Expert", "Alumni", "Parent", "Academician"}
    doc = docx.Document()
    doc.add_heading("Composition of BoS Members", level=1)
    cats = TEMPLATES["bos_composition"]["categories"]
    t = doc.add_table(rows=1, cols=4)
    for c, h in zip(t.rows[0].cells, ("S.No", "Category", "Role", "Members")):
        c.text = h
    for i, (name, _) in enumerate(cats, 1):
        who, what = people.get(name, (f"Dr. Sample Person {i}", name))
        role = roles.get(name, "External Invitee" if name in externals else
                         "Internal Invitee" if i > 5 else "Member")
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = str(i), name, role, f"{who}, {what}"
    doc.save(str(dst))
    return dst
