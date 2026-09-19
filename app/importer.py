"""
Department master importer.

Reads the Directors / Dy. Directors / Deans / HoDs contact workbook (or any
similar sheet) and maps its columns onto department records.  Column names
vary between versions of that file, so headers are matched by keyword rather
than by exact position, and the admin confirms the mapping before anything is
written.
"""

from __future__ import annotations

import io
import re

from openpyxl import load_workbook

# Candidate header keywords, in priority order, for each target field.
HEADER_HINTS = {
    "dept_name": ["department name", "department", "dept name", "dept"],
    "faculty": ["faculty"],
    "school": ["school", "college", "institute"],
    "dept_code": ["department code", "dept code", "code", "abbreviation", "abbr"],
    "campus": ["campus", "location", "city", "centre", "center"],
}

CAMPUS_ALIASES = {
    "bangalore": "Bengaluru", "bengaluru": "Bengaluru", "blr": "Bengaluru",
    "jain global campus": "Bengaluru", "jgi": "Bengaluru", "kanakapura": "Bengaluru",
    "jayanagar": "Bengaluru", "vv puram": "Bengaluru", "school street": "Bengaluru",
    "kochi": "Kochi", "cochin": "Kochi", "ernakulam": "Kochi", "kerala": "Kochi",
}


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9 ]+", " ", str(s or "").strip().lower()).strip()


def _score(header: str, hints: list[str]) -> int:
    h = _norm(header)
    if not h:
        return 0
    for i, hint in enumerate(hints):
        if h == hint:
            return 1000 - i
        if h.startswith(hint) or hint in h:
            return 500 - i
    return 0


def sniff_sheets(file_bytes: bytes):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    return wb.sheetnames


def find_header_row(rows, max_scan: int = 15):
    """The row with the most recognisable headers wins."""
    best, best_score = 0, -1
    for i, row in enumerate(rows[:max_scan]):
        score = sum(1 for cell in row if _norm(cell) and
                    any(_score(cell, hints) for hints in HEADER_HINTS.values()))
        if score > best_score:
            best, best_score = i, score
    return best


def build_mapping(headers):
    """Greedy best-match of each target field to a column index."""
    mapping, used = {}, set()
    for field, hints in HEADER_HINTS.items():
        best_idx, best_score = None, 0
        for idx, h in enumerate(headers):
            if idx in used:
                continue
            s = _score(h, hints)
            if s > best_score:
                best_idx, best_score = idx, s
        if best_idx is not None and best_score > 0:
            mapping[field] = best_idx
            used.add(best_idx)
    return mapping


def normalise_campus(value, default="Bengaluru"):
    v = _norm(value)
    if not v:
        return default
    for alias, canonical in CAMPUS_ALIASES.items():
        if alias in v:
            return canonical
    return value.strip().title() if isinstance(value, str) else default


def derive_code(dept_name: str, school: str = "", taken: set | None = None) -> str:
    """Make a short department code when the sheet does not supply one."""
    taken = taken or set()
    words = [w for w in re.split(r"[^A-Za-z0-9]+", dept_name or "") if w]
    stop = {"of", "and", "the", "for", "in", "department", "dept"}
    letters = "".join(w[0] for w in words if w.lower() not in stop).upper()[:6]
    if not letters:
        letters = re.sub(r"[^A-Z]", "", (dept_name or "DEPT").upper())[:4] or "DEPT"
    code, n = letters, 1
    while code in taken:
        n += 1
        code = f"{letters}{n}"
    return code


def parse_workbook(file_bytes: bytes, sheet_name: str | None = None,
                   mapping_override: dict | None = None,
                   default_campus: str = "Bengaluru"):
    """
    Returns (rows, meta).  `rows` are candidate department dicts;
    `meta` carries the detected headers and mapping so the admin can correct it.
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    grid = [[("" if c is None else str(c).strip()) for c in row]
            for row in ws.iter_rows(values_only=True)]
    grid = [r for r in grid if any(c for c in r)]
    if not grid:
        return [], {"error": "That sheet is empty.", "sheets": wb.sheetnames}

    hdr_i = find_header_row(grid)
    headers = grid[hdr_i]
    mapping = mapping_override or build_mapping(headers)

    rows, taken, seen_names = [], set(), set()
    for raw in grid[hdr_i + 1:]:
        def cell(field):
            idx = mapping.get(field)
            if idx is None or idx >= len(raw):
                return ""
            return str(raw[idx]).strip()

        name = cell("dept_name")
        if not name or _norm(name) in {"department", "total", "sl no", "s no"}:
            continue
        key = _norm(name) + "|" + _norm(cell("school"))
        if key in seen_names:
            continue
        seen_names.add(key)

        code = cell("dept_code") or derive_code(name, cell("school"), taken)
        code = re.sub(r"[^A-Za-z0-9\-/]", "", code).upper()[:20] or derive_code(name, "", taken)
        taken.add(code)

        rows.append({
            "dept_name": name,
            "faculty": cell("faculty"),
            "school": cell("school") or "—",
            "dept_code": code,
            "campus": normalise_campus(cell("campus"), default_campus),
        })

    meta = {
        "sheets": wb.sheetnames,
        "sheet": ws.title,
        "header_row": hdr_i + 1,
        "headers": headers,
        "mapping": mapping,
        "unmapped": [f for f in HEADER_HINTS if f not in mapping],
        "count": len(rows),
    }
    return rows, meta
