"""Rows for a form table, read from the department's own Excel or CSV file.

Curriculum matrices arrive in many shapes. This finds the header row by
looking for the table's column names (and the usual ways of writing them),
reads the rows under it, and turns each cell into what the column holds —
a number, a semester from "III" or "Semester 3", a course group matched to
the list. A "Semester II" line on its own sets the semester for the rows
below it, as many matrices are laid out that way.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

# other ways the columns are headed in real matrices
ALIASES = {
    "semester": ["sem", "semester", "semester no", "sem no"],
    "track": ["applies to", "track", "honours track"],
    "nep_category": ["course group", "category", "course category", "course type", "nep category",
                     "type of course", "component"],
    "course_code": ["code", "course code", "subject code", "paper code"],
    "course_title": ["title", "course title", "course name", "subject", "subject name",
                     "name of the course", "paper title", "course"],
    "l": ["l", "lecture", "lectures"],
    "t": ["t", "tutorial", "tutorials"],
    "p": ["p", "practical", "practicals", "lab"],
    "e": ["e", "experiential", "experiential learning"],
    "credits": ["credits", "credit", "cr", "c"],
    "cia": ["continuous assessment", "continuous assessment marks", "cia", "ia", "internal",
            "internal marks", "ca"],
    "ese": ["term end examination", "term end examination marks", "term end", "ese", "see",
            "external", "end semester", "end sem", "external marks"],
    "total_marks": ["total", "total marks", "max marks", "maximum marks", "marks"],
    "minor_title": ["minor", "minor stream", "stream"],
}

ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10}
SEM_LINE = re.compile(r"^\s*(?:semester|sem)\.?\s*[-:]?\s*([ivx]+|\d{1,2})\b", re.I)


class ImportError_(Exception):
    """Said to the person as it is."""


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def _cells(path: Path) -> list[list]:
    ext = path.suffix.lower()
    if ext == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return [row for row in csv.reader(io.StringIO(text))]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            raise ImportError_("This Excel file could not be opened — is it damaged or password-protected?")
        sheets = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        wb.close()
        return sheets          # a list of sheets; _best picks one
    if ext == ".xls":
        raise ImportError_("This is the old Excel format (.xls). Open it in Excel, choose "
                           "Save As → Excel Workbook (.xlsx), and upload that.")
    raise ImportError_("Only Excel (.xlsx) and CSV files can fill the table.")


def _match_headers(row: list, columns: list) -> dict[int, str]:
    """Which cell of this row heads which column."""
    found: dict[int, str] = {}
    taken = set()
    names = {c["name"]: [_norm(c["name"]), _norm(c.get("label"))] +
             [_norm(a) for a in ALIASES.get(c["name"], [])] for c in columns}
    # exact matches first, then a long alias inside a longer header
    for exact in (True, False):
        for i, cell in enumerate(row):
            h = _norm(cell)
            if not h or i in found:
                continue
            for name, alts in names.items():
                if name in taken:
                    continue
                if (exact and h in alts) or (not exact and any(len(a) >= 4 and a in h for a in alts)):
                    found[i] = name
                    taken.add(name)
                    break
    return found


def _best(sheets, columns):
    """The sheet and header row that match the most columns."""
    if sheets and sheets[0] and not isinstance(sheets[0][0], list):
        sheets = [sheets]      # a CSV: one sheet
    best = (0, None, None, None)
    for rows in sheets:
        for r, row in enumerate(rows[:40]):
            m = _match_headers(row or [], columns)
            if len(m) > best[0]:
                best = (len(m), rows, r, m)
    return best


def _int(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip()
    if s.lower() in ROMAN:
        return ROMAN[s.lower()]
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return int(round(float(m.group()))) if m else None


def _semester(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    m = SEM_LINE.match(s) or re.match(r"^([ivx]+|\d{1,2})$", s)
    if m:
        t = m.group(1)
        return ROMAN.get(t) or (int(t) if t.isdigit() else None)
    return _int(v)


def _option(v, options):
    s = _norm(v)
    if not s:
        return ""
    for o in options:
        if _norm(o) == s:
            return o
    for o in options:                    # "Major" → "Major (Core)", "AEC" → "Ability Enhancement …"
        no = _norm(o)
        if no.startswith(s) or s in no.split() or (len(s) >= 4 and s in no):
            return o
    return str(v).strip()                # kept as written; the check flags it


def rows_from_file(path: Path, section: dict) -> dict:
    columns = [c for c in section.get("columns", []) if not c.get("auto_index")]
    n, rows, head, matched = _best(_cells(path), columns)
    if n < 2:
        raise ImportError_("No header row found that matches this table. The first row of the "
                           "table in your file should have headings such as Semester, Course code, "
                           "Course title, L, T, P, Credits.")
    by_name = {c["name"]: c for c in columns}
    out, skipped, current_sem = [], 0, None
    for raw in rows[head + 1:]:
        raw = list(raw or [])
        texts = [str(c).strip() for c in raw if c is not None and str(c).strip() != ""]
        if not texts:
            continue
        # a "Semester III" line on its own sets the semester below it
        if len(texts) <= 2 and SEM_LINE.match(texts[0]):
            current_sem = _semester(texts[0])
            continue
        row = {}
        for i, name in matched.items():
            v = raw[i] if i < len(raw) else None
            col = by_name[name]
            if name == "semester":
                v = _semester(v)
            elif col.get("type") == "integer":
                v = _int(v)
            elif col.get("type") == "select":
                v = _option(v, col.get("options", []))
            else:
                v = "" if v is None else re.sub(r"\s+", " ", str(v)).strip()
            if v not in (None, ""):
                row[name] = v
        if "semester" in by_name and "semester" not in row and current_sem:
            row["semester"] = current_sem
        title = str(row.get("course_title") or row.get("minor_title") or "")
        if not (title or row.get("course_code")) or re.match(r"^\s*(grand\s+)?total\b", title, re.I):
            skipped += 1
            continue
        out.append(row)
    if not out:
        raise ImportError_("The headings were found, but no course rows under them.")
    return {"rows": out, "skipped": skipped,
            "columns": sorted(set(matched.values()), key=list(by_name).index),
            "missing": [by_name[k]["label"] for k in by_name
                        if by_name[k].get("required") and k not in matched.values()
                        and not all(k in r for r in out)]}
